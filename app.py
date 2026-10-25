import streamlit as st
import pandas as pd
import os
import io
import re
import math
import json
import time
import tempfile
import subprocess
from datetime import datetime
from PIL import Image, ImageOps
try:
    from PIL import ImageGrab
except ImportError:
    ImageGrab = None

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Configuration
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PARTS = os.path.join(BASE_DIR, "FST16_DT_Parts_Export.xlsx")
EXCEL_ASSEMBLIES = os.path.join(BASE_DIR, "FST16_Assemblies_Export_DL.xlsx")
EXCEL_TRASH = os.path.join(BASE_DIR, "FST16_Trash_Internal.xlsx")
EXCEL_REQUIREMENTS = os.path.join(BASE_DIR, "FST16_Requirements.xlsx")
IMAGE_DIR = os.path.join(BASE_DIR, "images")

if not os.path.exists(IMAGE_DIR):
    os.makedirs(IMAGE_DIR)

# ==================== IMAGE UTILITIES ====================
def resolve_image_path(path_str):
    if not path_str or pd.isna(path_str):
        return None
    clean_path = str(path_str).replace("\\", "/").strip()
    if os.path.isabs(clean_path) and os.path.exists(clean_path):
        return clean_path
    full_path = os.path.join(BASE_DIR, clean_path)
    if os.path.exists(full_path):
        return full_path
    if os.path.exists(clean_path):
        return clean_path
    return None

def grab_image_from_clipboard():
    if ImageGrab is None:
        return None
    for _ in range(3):
        try:
            img = ImageGrab.grabclipboard()
            if isinstance(img, Image.Image):
                return img
            elif isinstance(img, list) and len(img) > 0:
                if os.path.exists(img[0]):
                    return Image.open(img[0])
        except Exception:
            pass
        time.sleep(0.1)
    return None

def extract_ntop_fea_from_image(pil_img, target_part=None):
    if pil_img is None or os.name != "nt":
        return None, None
    tmp_path = None
    try:
        im = pil_img.convert("L")
        w, h = im.size
        crop = ImageOps.invert(im.crop((int(w * 0.68), int(h * 0.22), w, h)))
        crop = crop.resize((crop.width * 3, crop.height * 3), Image.Resampling.BICUBIC)
        fd, tmp_path = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        crop.save(tmp_path, format="PNG")

        ps_cmd = (
            "Add-Type -AssemblyName System.Runtime.WindowsRuntime; "
            "$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType = WindowsRuntime]; "
            "$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Foundation, ContentType = WindowsRuntime]; "
            "$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType = WindowsRuntime]; "
            "$g = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]; "
            "Function Await($t, $r) { $a = $g.MakeGenericMethod($r).Invoke($null, @($t)); $a.Wait(-1) | Out-Null; $a.Result }; "
            "$eng = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages(); "
            f"$f = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync('{tmp_path}')) ([Windows.Storage.StorageFile]); "
            "$s = Await ($f.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream]); "
            "$d = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($s)) ([Windows.Graphics.Imaging.BitmapDecoder]); "
            "$b = Await ($d.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap]); "
            "$res = Await ($eng.RecognizeAsync($b)) ([Windows.Media.Ocr.OcrResult]); "
            "$s.Dispose(); "
            "foreach ($l in $res.Lines) { Write-Output $l.Text }"
        )
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps_cmd],
            capture_output=True,
            text=True,
            timeout=5
        )
        text = proc.stdout or ""
        if not text.strip():
            return None, None

        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        first_lines = " ".join(lines[:4]).lower()
        nospace = first_lines.replace(" ", "")

        detected_lc = None
        # Load Case detection: support UP_01 and HB_01 matchers
        if "pure" in first_lines and "brak" in first_lines:
            detected_lc = "LC1"
        elif "pure" in first_lines and "acc" in first_lines:
            detected_lc = "LC2"
        elif "pure" in first_lines and ("corn" in first_lines or "com" in first_lines):
            detected_lc = "LC3"
        elif "b+c2" in nospace or "b+c_2" in nospace or ("brak" in first_lines and ("cornering2" in nospace or "comering2" in nospace)):
            detected_lc = "LC5"
        elif "b+c1" in nospace or "b+c_1" in nospace or "b+c" in nospace or ("brak" in first_lines and ("corn" in first_lines or "com" in first_lines)):
            detected_lc = "LC4"
        elif "a+c2" in nospace or "a+c_2" in nospace or ("acc" in first_lines and ("cornering2" in nospace or "comering2" in nospace)):
            detected_lc = "LC7"
        elif "a+c1" in nospace or "a+c_1" in nospace or "a+c" in nospace or ("acc" in first_lines and ("corn" in first_lines or "com" in first_lines)):
            detected_lc = "LC6"
        elif target_part != "HB_01":
            if "-1kn" in nospace or "-ikn" in nospace or "result_2" in nospace or "result-2" in nospace:
                detected_lc = "LC9"
            elif "1kn" in nospace or "ikn" in nospace:
                detected_lc = "LC8"

        nums = []
        for m in re.finditer(r"([1-9])[.\-]?(\d{5})\s*(?:e\s*[\-+]|[\+&])\s*0([0-4])", text, re.I):
            mant = float(m.group(1) + "." + m.group(2))
            exp = int(m.group(3))
            nums.append(mant * (10 ** (-exp)))

        detected_disp = None
        if nums:
            raw_max = nums[0]
            part_decimals = 5 if target_part == "HB_01" else 4
            factor = 10 ** part_decimals
            detected_disp = math.floor(raw_max * factor + 1e-9) / factor

        return detected_lc, detected_disp
    except Exception:
        return None, None
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

def render_image_input(label="Schematic", key="img_input", current_image_path=None, ocr_callback=None):
    clipboard_key = f"clipboard_{key}"
    st.markdown(f"**{label}**")

    col_paste, col_clear = st.columns([1, 1])
    with col_paste:
        pasted_img = st.session_state.get(clipboard_key, None)
        btn_label = "Replace Clipboard" if (pasted_img is not None or current_image_path) else "Paste Clipboard"
        if st.button(btn_label, key=f"btn_paste_{key}", use_container_width=True):
            pasted = grab_image_from_clipboard()
            if pasted is not None:
                st.session_state[clipboard_key] = pasted
                if ocr_callback is not None:
                    ocr_callback(pasted)
                else:
                    st.toast("Image pasted from clipboard.")
                st.rerun()
            else:
                st.warning("No image found on system clipboard. Capture a screenshot first (Win + Shift + S).")

    with col_clear:
        pasted_img = st.session_state.get(clipboard_key, None)
        if pasted_img is not None:
            if st.button("Clear Image", key=f"btn_clear_{key}", use_container_width=True):
                st.session_state.pop(clipboard_key, None)
                st.rerun()

    pasted_img = st.session_state.get(clipboard_key, None)
    if pasted_img is not None:
        st.image(pasted_img, caption="New Image", use_container_width=True)
    elif current_image_path:
        real_img = resolve_image_path(current_image_path)
        if real_img:
            st.image(real_img, caption="Current Image", use_container_width=True)

    return pasted_img

def save_pasted_image(pasted_image, part_name, version_str):
    if pasted_image is None:
        return ""
    clean_name = str(part_name).replace(" ", "_")
    version_clean = str(version_str).lstrip('vV').replace('.', '_').replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{clean_name}_v{version_clean}_{timestamp}.png"
    filepath = os.path.join(IMAGE_DIR, filename)
    pasted_image.save(filepath, format="PNG")
    return f"images/{filename}"

def parse_lc_images(val):
    if pd.isna(val) or str(val).strip() in ["", "nan", "None"]:
        return {}
    try:
        data = json.loads(str(val).strip())
        if isinstance(data, dict):
            return data
    except Exception:
        pass
    return {}

def format_lc_images(img_dict):
    cleaned = {k: v for k, v in img_dict.items() if v and str(v).strip()}
    return json.dumps(cleaned)

def save_lc_image(pasted_image, part_name, version_str, lc_key):
    if pasted_image is None:
        return ""
    clean_name = str(part_name).replace(" ", "_")
    version_clean = str(version_str).lstrip('vV').replace('.', '_').replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{clean_name}_v{version_clean}_{lc_key}_{timestamp}.png"
    filepath = os.path.join(IMAGE_DIR, filename)
    pasted_image.save(filepath, format="PNG")
    return f"images/{filename}"

# ==================== TEXT & FORMATTING UTILITIES ====================
def calc_description_height(text, min_height=100, max_height=600):
    if not text:
        return min_height
    text_str = str(text)
    lines = text_str.count('\n') + 1
    wrapped_lines = sum(max(1, len(line) // 65) for line in text_str.split('\n'))
    total_lines = max(lines, wrapped_lines)
    return max(min_height, min(max_height, total_lines * 24 + 40))

def format_v_label(v):
    v_str = str(v)
    return v_str if v_str.startswith("v") or v_str.startswith("FST") or v_str.startswith("OEM") else f"v{v_str}"

def parse_measurements(val):
    if pd.isna(val) or str(val).strip() == "":
        return []
    lines = str(val).split("\n")
    res = []
    for l in lines:
        if ":" in l:
            k, v = l.split(":", 1)
            res.append({"Property": k.strip(), "Value": v.strip()})
        else:
            if l.strip():
                res.append({"Property": "Note", "Value": l.strip()})
    return res

def format_measurements(records):
    lines = [f"{str(r.get('Property', '')).strip()}: {str(r.get('Value', '')).strip()}" 
             for r in records if str(r.get('Property', '')).strip() != ""]
    return "\n".join(lines)

def parse_config(val):
    if pd.isna(val) or str(val).strip() == "":
        return {}
    res = {}
    for l in str(val).split("\n"):
        if ":" in l:
            k, v = l.split(":", 1)
            clean_v = v.strip()
            if clean_v.startswith("v") or clean_v.startswith("V"):
                clean_v = clean_v[1:]
            res[k.strip()] = clean_v
    return res

def format_config(config_dict):
    return "\n".join([f"{k}: v{v}" for k, v in config_dict.items()])

def parse_tolerancing_data(val):
    if val is None:
        return []
    if isinstance(val, list):
        data = val
    elif pd.isna(val) or str(val).strip() in ["", "nan", "None"]:
        return []
    else:
        try:
            data = json.loads(str(val).strip())
        except Exception:
            return []
    if isinstance(data, list):
        normalized = []
        for item in data:
            if isinstance(item, dict):
                feat = str(item.get("Interface / Feature", item.get("Interface / Elemento", ""))).strip()
                if not feat or feat.lower() in ["none", "nan"]:
                    continue
                norm_item = {
                    "Interface / Feature": feat,
                    "Nominal (mm)": str(item.get("Nominal (mm)", "")).strip(),
                    "Fit / Tolerance": str(item.get("Fit / Tolerance", item.get("Ajuste / Tolerância", ""))).strip(),
                    "Roughness / Notes": str(item.get("Roughness / Notes", item.get("Rugosidade / Notas", ""))).strip()
                }
                normalized.append(norm_item)
        return normalized
    return []

def format_tolerancing_data(records):
    cleaned = []
    for r in records:
        if isinstance(r, dict) and any(str(v).strip() for v in r.values()):
            cleaned.append({
                "Interface / Feature": str(r.get("Interface / Feature", r.get("Interface / Elemento", ""))).strip(),
                "Nominal (mm)": str(r.get("Nominal (mm)", "")).strip(),
                "Fit / Tolerance": str(r.get("Fit / Tolerance", r.get("Ajuste / Tolerância", ""))).strip(),
                "Roughness / Notes": str(r.get("Roughness / Notes", r.get("Rugosidade / Notas", ""))).strip()
            })
    return json.dumps(cleaned)

# ==================== EXCEL TABLE FORMATTER ====================
def read_excel_safely(filepath):
    try:
        with open(filepath, "rb") as f:
            buf = io.BytesIO(f.read())
        return pd.read_excel(buf)
    except PermissionError:
        if os.name == "nt":
            import ctypes
            import msvcrt
            GENERIC_READ = 0x80000000
            FILE_SHARE_ALL = 0x00000001 | 0x00000002 | 0x00000004
            OPEN_EXISTING = 3
            FILE_ATTRIBUTE_NORMAL = 0x80
            handle = ctypes.windll.kernel32.CreateFileW(
                str(filepath), GENERIC_READ, FILE_SHARE_ALL, None, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL, None
            )
            if handle not in (-1, 0xFFFFFFFF, 0xFFFFFFFFFFFFFFFF):
                fd = msvcrt.open_osfhandle(handle, os.O_RDONLY | os.O_BINARY)
                with os.fdopen(fd, "rb") as f:
                    buf = io.BytesIO(f.read())
                return pd.read_excel(buf)
        raise

def save_formatted_excel(df, filename, table_name="DataTable"):
    try:
        if "Subassembly" in df.columns and "Part Name" in df.columns:
            df = df.sort_values(by=["Subassembly", "Part Name"]).reset_index(drop=True)
            
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            sheet_name = 'Data'
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            
            for idx, col in enumerate(df.columns, 1):
                col_letter = get_column_letter(idx)
                
                if col in ["Description", "Specifications", "Constraints", "Part Configuration", "Image Path", 
                            "Stiffness Data", "LC Images", "Tolerancing Data", "Justification",
                            "Requirement", "Explanation", "Validation", "Related Parts"]:
                    worksheet.column_dimensions[col_letter].width = 45
                    for row in range(2, len(df) + 2):
                        worksheet.cell(row=row, column=idx).alignment = Alignment(wrap_text=True, vertical='top')
                else:
                    col_lens = [len(str(x)) for x in df[col].dropna()]
                    max_len = max(max(col_lens) if col_lens else 0, len(str(col)))
                    worksheet.column_dimensions[col_letter].width = max(min(max_len + 5, 25), 14)
                    for row in range(2, len(df) + 2):
                        worksheet.cell(row=row, column=idx).alignment = Alignment(horizontal="center", vertical='top')
                        
            if not df.empty:
                max_row = len(df) + 1
                max_col = len(df.columns)
                clean_table_name = table_name.replace(" ", "_").replace("-", "_")
                tab = Table(displayName=clean_table_name, ref=f"A1:{get_column_letter(max_col)}{max_row}")
                style = TableStyleInfo(
                    name="TableStyleMedium2", showFirstColumn=False,
                    showLastColumn=False, showRowStripes=True, showColumnStripes=False
                )
                tab.tableStyleInfo = style
                worksheet.add_table(tab)
        return True
    except PermissionError:
        st.error(f"Close '{os.path.basename(filename)}' in Microsoft Excel before saving.")
        return False

# ==================== FEA STIFFNESS & DISPLACEMENTS ====================
FEA_PARTS_CONFIG = {
    "UP_01": {
        "DECIMALS": 4,
        "LC_KEYS": ["LC1", "LC2", "LC3", "LC4", "LC5", "LC6", "LC7", "LC8", "LC9"],
        "LC_INFO": {
            "LC1": {"desc": "Pure Braking", "ref_fst15": 0.2955, "default_fst16": 0.3211},
            "LC2": {"desc": "Pure Acceleration", "ref_fst15": 0.4246, "default_fst16": 0.4234},
            "LC3": {"desc": "Pure Cornering", "ref_fst15": 0.3381, "default_fst16": 0.5343},
            "LC4": {"desc": "Braking + Cornering 1", "ref_fst15": 0.4214, "default_fst16": 0.4075},
            "LC5": {"desc": "Braking + Cornering 2", "ref_fst15": 0.4229, "default_fst16": 0.3582},
            "LC6": {"desc": "Acceleration + Cornering 1", "ref_fst15": 0.3262, "default_fst16": 0.4746},
            "LC7": {"desc": "Acceleration + Cornering 2", "ref_fst15": 0.3574, "default_fst16": 0.5312},
            "LC8": {"desc": "1kN Fy", "ref_fst15": 0.0, "default_fst16": 0.0},
            "LC9": {"desc": "-1kN Fy", "ref_fst15": 0.0, "default_fst16": 0.0},
        },
        "DEFAULT_FST15": {
            "LC1": 0.2955, "LC2": 0.4246, "LC3": 0.3381, "LC4": 0.4214,
            "LC5": 0.4229, "LC6": 0.3262, "LC7": 0.3574, "LC8": 0.0, "LC9": 0.0
        },
        "DEFAULT_FST16": {
            "LC1": 0.3211, "LC2": 0.4234, "LC3": 0.5343, "LC4": 0.4075,
            "LC5": 0.3582, "LC6": 0.4746, "LC7": 0.5312, "LC8": 0.0, "LC9": 0.0
        }
    },
    "HB_01": {
        "DECIMALS": 5,
        "LC_KEYS": ["LC1", "LC2", "LC3", "LC4", "LC5", "LC6", "LC7", "LC8", "LC9"],
        "LC_INFO": {
            "LC1": {"desc": "Pure Braking", "ref_fst15": 0.0, "default_fst16": None},
            "LC2": {"desc": "Pure Acceleration", "ref_fst15": 0.0, "default_fst16": None},
            "LC3": {"desc": "Pure Cornering", "ref_fst15": 0.0, "default_fst16": None},
            "LC4": {"desc": "Braking + Cornering 1", "ref_fst15": 0.0, "default_fst16": None},
            "LC5": {"desc": "Braking + Cornering 2", "ref_fst15": 0.0, "default_fst16": None},
            "LC6": {"desc": "Acceleration + Cornering 1", "ref_fst15": 0.0, "default_fst16": None},
            "LC7": {"desc": "Acceleration + Cornering 2", "ref_fst15": 0.0, "default_fst16": None},
            "LC8": {"desc": "1kN Fy", "ref_fst15": 0.0, "default_fst16": None},
            "LC9": {"desc": "-1kN Fy", "ref_fst15": 0.0, "default_fst16": None},
        },
        "DEFAULT_FST15": {
            "LC1": 0.0, "LC2": 0.0, "LC3": 0.0, "LC4": 0.0,
            "LC5": 0.0, "LC6": 0.0, "LC7": 0.0, "LC8": 0.0, "LC9": 0.0
        },
        "DEFAULT_FST16": {
            "LC1": 0.0, "LC2": 0.0, "LC3": 0.0, "LC4": 0.0,
            "LC5": 0.0, "LC6": 0.0, "LC7": 0.0, "LC8": 0.0, "LC9": 0.0
        }
    }
}

def get_part_fea_config(part_name):
    clean = str(part_name).strip()
    return FEA_PARTS_CONFIG.get(clean, {"DECIMALS": 4, "LC_KEYS": [], "LC_INFO": {}, "DEFAULT_FST15": {}, "DEFAULT_FST16": {}})

# Backwards compatibility globals
LC_KEYS = FEA_PARTS_CONFIG["UP_01"]["LC_KEYS"]
LC_INFO = FEA_PARTS_CONFIG["UP_01"]["LC_INFO"]
DEFAULT_FST15_DISPLACEMENTS = FEA_PARTS_CONFIG["UP_01"]["DEFAULT_FST15"]
DEFAULT_FST16_DISPLACEMENTS = FEA_PARTS_CONFIG["UP_01"]["DEFAULT_FST16"]

def parse_stiffness_data(val, lc_keys=None):
    if pd.isna(val) or str(val).strip() == "" or str(val).strip() == "nan":
        return {}
    val_str = str(val).strip().replace(": NaN", ": 0.0").replace(": nan", ": 0.0")
    try:
        raw = json.loads(val_str)
        if not isinstance(raw, dict):
            return {}
        res = {}
        target_keys = lc_keys if lc_keys is not None else list(raw.keys())
        for lc in target_keys:
            if lc in raw:
                item = raw[lc]
                if isinstance(item, dict):
                    num = float(item.get("disp", 0.0) or 0.0)
                else:
                    num = float(item if item is not None else 0.0)
                res[lc] = 0.0 if math.isnan(num) else num
        return res
    except Exception:
        pass
    return {}

def format_stiffness_data(disp_dict, lc_keys=None, decimals=4):
    cleaned = {}
    target_keys = lc_keys if lc_keys is not None else list(disp_dict.keys())
    for lc in target_keys:
        if lc in disp_dict and disp_dict[lc] is not None:
            try:
                if isinstance(disp_dict[lc], dict):
                    num = float(disp_dict[lc].get("disp", 0.0) or 0.0)
                else:
                    num = float(disp_dict[lc])
                cleaned[lc] = 0.0 if math.isnan(num) else round(num, decimals)
            except (ValueError, TypeError):
                cleaned[lc] = 0.0
    return json.dumps(cleaned)

def build_comparison_dataframe(curr_stiff, ref_stiff, ref_label="FST15", curr_label=None, lc_info=None, lc_keys=None):
    rows = []
    curr_col = f"Displacement ({curr_label}) (mm)" if curr_label else "Displacement (mm)"
    ref_col = f"Ref. ({ref_label}) (mm)"
    target_keys = lc_keys if lc_keys is not None else LC_KEYS
    target_info = lc_info if lc_info is not None else LC_INFO
    for lc in target_keys:
        desc = target_info.get(lc, {}).get("desc", lc)
        
        c_item = curr_stiff.get(lc, 0.0)
        if c_item is None:
            c_disp = 0.0
        elif isinstance(c_item, dict):
            c_disp = float(c_item.get("disp", 0.0) or 0.0)
        else:
            try:
                c_disp = float(c_item)
            except (ValueError, TypeError):
                c_disp = 0.0

        r_item = ref_stiff.get(lc, 0.0)
        if r_item is None:
            r_disp = 0.0
        elif isinstance(r_item, dict):
            r_disp = float(r_item.get("disp", 0.0) or 0.0)
        else:
            try:
                r_disp = float(r_item)
            except (ValueError, TypeError):
                r_disp = 0.0
        
        diff = c_disp - r_disp
        if abs(diff) < 1e-7:
            diff = 0.0
            pct = 0.0
        else:
            pct = (diff / r_disp * 100.0) if r_disp != 0 else 0.0
            
        rows.append({
            "Load Case": lc,
            "Description": desc,
            curr_col: c_disp,
            ref_col: r_disp,
            "Diff (mm)": diff,
            "%": pct
        })
    return pd.DataFrame(rows)

def style_comparison_dataframe(df_comp, decimals=4):
    def highlight_delta(val):
        try:
            v = float(val)
            if v < -0.000001:
                return 'color: #33cc33; font-weight: 600;'
            elif v > 0.000001:
                return 'color: #ff4b4b; font-weight: 600;'
        except Exception:
            pass
        return ''

    delta_col = "Diff (mm)"
    pct_col = "%"

    format_dict = {}
    for c in df_comp.columns:
        if c == delta_col:
            format_dict[c] = f"{{:+.{decimals}f}}"
        elif c == pct_col:
            format_dict[c] = "{:+.1f}%"
        elif "(mm)" in c:
            format_dict[c] = f"{{:.{decimals}f}}"

    styled = df_comp.style.map(
        highlight_delta,
        subset=[c for c in df_comp.columns if c in [delta_col, pct_col]]
    ).format(format_dict)
    return styled

def build_mass_comparison_dataframe(eval_weight, ref_weight, curr_label="Current", ref_label="FST15"):
    try:
        e_w = float(eval_weight)
    except Exception:
        e_w = 0.0
    try:
        r_w = float(ref_weight)
    except Exception:
        r_w = 0.0

    diff = e_w - r_w
    if abs(diff) < 1e-4:
        diff = 0.0
        pct = 0.0
    else:
        pct = (diff / r_w * 100.0) if r_w != 0 else 0.0

    curr_col = f"Weight ({curr_label}) (g)" if curr_label else "Weight (g)"
    ref_col = f"Ref. ({ref_label}) (g)"

    return pd.DataFrame([{
        "Metric": "Total Weight",
        curr_col: e_w,
        ref_col: r_w,
        "Diff (g)": diff,
        "%": pct
    }])

def build_all_iterations_mass_dataframe(part_data, ref_weight, ref_label="FST15"):
    try:
        r_w = float(ref_weight)
    except Exception:
        r_w = 0.0

    ref_col = f"Ref. ({ref_label}) (g)"
    rows = []
    for _, r in part_data.sort_values("Date").iterrows():
        try:
            w = float(r["Weight (g)"])
        except Exception:
            w = 0.0
        diff = w - r_w
        if abs(diff) < 1e-4:
            diff = 0.0
            pct = 0.0
        else:
            pct = (diff / r_w * 100.0) if r_w != 0 else 0.0

        rows.append({
            "Version": format_v_label(r["Version"]),
            "Date": str(r["Date"]),
            "Weight (g)": w,
            ref_col: r_w,
            "Diff (g)": diff,
            "%": pct
        })
    return pd.DataFrame(rows)

def style_mass_dataframe(df_mass):
    def highlight_delta(val):
        try:
            v = float(val)
            if v < -0.01:
                return 'color: #33cc33; font-weight: 600;'
            elif v > 0.01:
                return 'color: #ff4b4b; font-weight: 600;'
        except Exception:
            pass
        return ''

    delta_col = "Diff (g)"
    pct_col = "%"

    format_dict = {}
    for c in df_mass.columns:
        if c == delta_col:
            format_dict[c] = "{:+.2f}"
        elif c == pct_col:
            format_dict[c] = "{:+.1f}%"
        elif "(g)" in c:
            format_dict[c] = "{:.2f}"

    styled = df_mass.style.map(
        highlight_delta,
        subset=[c for c in df_mass.columns if c in [delta_col, pct_col]]
    ).format(format_dict)
    return styled

def parse_structural_data(val):
    if pd.isna(val) or str(val).strip() in ["", "nan", "None"]:
        return {}
    try:
        data = json.loads(str(val).strip())
        if isinstance(data, dict):
            if "Structural" in data and isinstance(data["Structural"], dict):
                return data["Structural"]
            if "stress" in data or "peak_stress" in data or "sf" in data:
                return data
    except Exception:
        pass
    return {}

def format_structural_data(stress, sf, disp, lc_name="Bump / Lateral", yield_strength=503.0, material="Al 7075-T6"):
    try:
        s = float(stress)
    except Exception:
        s = 0.0
    try:
        y = float(yield_strength)
    except Exception:
        y = 503.0
    try:
        d = float(disp)
    except Exception:
        d = 0.0
    calculated_sf = round(y / s, 2) if s > 0 else float(sf or 0.0)
    data = {
        "load_case": str(lc_name).strip() if lc_name else "Bump / Lateral",
        "stress": s,
        "yield_strength": y,
        "sf": calculated_sf,
        "disp": d,
        "material": str(material).strip() if material else "Al 7075-T6"
    }
    return json.dumps(data)

def build_bracket_comparison_dataframe(eval_data, ref_data, eval_label="Current", ref_label="Baseline"):
    e_stress = float(eval_data.get("stress", 0.0))
    r_stress = float(ref_data.get("stress", 0.0))
    diff_stress = e_stress - r_stress

    e_sf = float(eval_data.get("sf", 0.0))
    r_sf = float(ref_data.get("sf", 0.0))
    diff_sf = e_sf - r_sf

    e_disp = float(eval_data.get("disp", 0.0))
    r_disp = float(ref_data.get("disp", 0.0))
    diff_disp = e_disp - r_disp

    eval_col = f"Eval ({eval_label})"
    ref_col = f"Ref ({ref_label})"

    rows = [
        {
            "Metric": "Critical Load Case",
            eval_col: str(eval_data.get("load_case", "-")),
            ref_col: str(ref_data.get("load_case", "-")),
            "Delta": "-",
            "Criteria / Target": "-"
        },
        {
            "Metric": "Material & Yield Strength",
            eval_col: f"{eval_data.get('material', 'Al 7075-T6')} ({eval_data.get('yield_strength', 503.0):.0f} MPa)",
            ref_col: f"{ref_data.get('material', 'Al 7075-T6')} ({ref_data.get('yield_strength', 503.0):.0f} MPa)",
            "Delta": "-",
            "Criteria / Target": "Sy >= 503 MPa"
        },
        {
            "Metric": "Peak von Mises Stress (MPa)",
            eval_col: f"{e_stress:.1f}" if e_stress > 0 else "-",
            ref_col: f"{r_stress:.1f}" if r_stress > 0 else "-",
            "Delta": f"{diff_stress:+.1f} MPa" if (e_stress > 0 and r_stress > 0) else "-",
            "Criteria / Target": f"Limit: {float(eval_data.get('yield_strength', 503.0)):.1f} MPa"
        },
        {
            "Metric": "Safety Factor (SF)",
            eval_col: f"{e_sf:.2f}" if e_sf > 0 else "-",
            ref_col: f"{r_sf:.2f}" if r_sf > 0 else "-",
            "Delta": f"{diff_sf:+.2f}" if (e_sf > 0 and r_sf > 0) else "-",
            "Criteria / Target": "Target SF >= 2.03"
        },
        {
            "Metric": "Max Deflection (mm)",
            eval_col: f"{e_disp:.4f}" if e_disp > 0 else "-",
            ref_col: f"{r_disp:.4f}" if r_disp > 0 else "-",
            "Delta": f"{diff_disp:+.4f} mm" if (e_disp > 0 and r_disp > 0) else "-",
            "Criteria / Target": "Stiffness compliance"
        }
    ]
    return pd.DataFrame(rows)

# ==================== DEFAULT KNOWLEDGE BASE ====================
DEFAULT_JUSTIFICATIONS = {
    "UP_03": "Motor Front Plate providing machined sealing face for motor cooling jacket O-ring and reaction keys for motor stator torque against the upright. Retained as separate plate to guarantee seal integrity and quick motor swaps.",
    "UP_04": "Motor Back Plate providing an accessible visual installation plane for the motor cooling jacket O-ring, eliminating blind insertion leak risks during motor assembly, and shielding the upright cavity from thermal soak."
}

DEFAULT_OEM_TOLERANCES = {}

DEFAULT_CUSTOM_TOLERANCES = {}

# ==================== DATA LOAD & SAVE ====================
def load_parts():
    if os.path.exists(EXCEL_PARTS):
        try:
            df = read_excel_safely(EXCEL_PARTS)
            text_cols = [
                "Subassembly", "Part Name", "Version", "Date", 
                "Description", "Specifications", "Image Path", 
                "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
            ]
            for c in text_cols:
                if c not in df.columns:
                    df[c] = ""
                else:
                    df[c] = df[c].fillna("").astype(object)
            if "Component Status" in df.columns:
                df = df.drop(columns=["Component Status"])
                needs_save = True
            else:
                needs_save = False

            # Deduplicate any duplicate (Part Name, Version) rows keeping the latest entry
            before_len = len(df)
            df = df.drop_duplicates(subset=["Part Name", "Version"], keep="last").reset_index(drop=True)
            if len(df) != before_len:
                needs_save = True

            # Auto-seed and sanitize UP_01 stiffness data
            up01_mask = df["Part Name"].astype(str) == "UP_01"
            if up01_mask.any():
                for idx in df[up01_mask].index.tolist():
                    val = df.at[idx, "Stiffness Data"]
                    val_str = str(val).strip()
                    if pd.isna(val) or val_str in ["", "nan"]:
                        ver = str(df.at[idx, "Version"])
                        if "FST15" in ver.upper():
                            df.at[idx, "Stiffness Data"] = format_stiffness_data(DEFAULT_FST15_DISPLACEMENTS)
                            needs_save = True
                        else:
                            df.at[idx, "Stiffness Data"] = format_stiffness_data(DEFAULT_FST16_DISPLACEMENTS)
                            needs_save = True
                    elif "NaN" in val_str or "nan" in val_str:
                        df.at[idx, "Stiffness Data"] = format_stiffness_data(parse_stiffness_data(val_str))
                        needs_save = True

            if needs_save:
                try:
                    save_parts(df)
                except Exception:
                    pass

            return df
        except PermissionError:
            st.error(f"Close '{os.path.basename(EXCEL_PARTS)}' in Microsoft Excel to continue.")
            return pd.DataFrame(columns=[
                "Subassembly", "Part Name", "Version", "Date", 
                "Weight (g)", "Description", "Specifications", "Image Path", 
                "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
            ])
    else:
        return pd.DataFrame(columns=[
            "Subassembly", "Part Name", "Version", "Date", 
            "Weight (g)", "Description", "Specifications", "Image Path", 
            "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
        ])

def save_parts(df):
    if "Component Status" in df.columns:
        df = df.drop(columns=["Component Status"])
    return save_formatted_excel(df, EXCEL_PARTS, table_name="PartsTable")

def load_assemblies():
    if os.path.exists(EXCEL_ASSEMBLIES):
        try:
            df = read_excel_safely(EXCEL_ASSEMBLIES)
            for c in ["Assembly Version", "Date", "Description", "Part Configuration"]:
                if c in df.columns:
                    df[c] = df[c].fillna("").astype(object)
                else:
                    df[c] = ""
            return df
        except PermissionError:
            st.error(f"Close '{os.path.basename(EXCEL_ASSEMBLIES)}' in Microsoft Excel to continue.")
            return pd.DataFrame(columns=[
                "Assembly Version", "Date", "Total Weight (g)", 
                "Description", "Part Configuration"
            ])
    else:
        return pd.DataFrame(columns=[
            "Assembly Version", "Date", "Total Weight (g)", 
            "Description", "Part Configuration"
        ])

def save_assemblies(df):
    return save_formatted_excel(df, EXCEL_ASSEMBLIES, table_name="AssembliesTable")

def load_trash():
    if os.path.exists(EXCEL_TRASH):
        try:
            df = read_excel_safely(EXCEL_TRASH)
            if "Component Status" in df.columns:
                df = df.drop(columns=["Component Status"])
            text_cols = [
                "Subassembly", "Part Name", "Version", "Date", 
                "Description", "Specifications", "Image Path", 
                "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
            ]
            for c in text_cols:
                if c not in df.columns:
                    df[c] = ""
                else:
                    df[c] = df[c].fillna("").astype(object)
            return df
        except PermissionError:
            st.error(f"Close '{os.path.basename(EXCEL_TRASH)}' in Microsoft Excel to continue.")
            return pd.DataFrame(columns=[
                "Subassembly", "Part Name", "Version", "Date", 
                "Weight (g)", "Description", "Specifications", "Image Path", 
                "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
            ])
    else:
        return pd.DataFrame(columns=[
            "Subassembly", "Part Name", "Version", "Date", 
            "Weight (g)", "Description", "Specifications", "Image Path", 
            "Stiffness Data", "LC Images", "Tolerancing Data", "Justification"
        ])

def save_trash(df):
    if "Component Status" in df.columns:
        df = df.drop(columns=["Component Status"])
    return save_formatted_excel(df, EXCEL_TRASH, table_name="TrashTable")

def load_requirements():
    if os.path.exists(EXCEL_REQUIREMENTS):
        try:
            df = read_excel_safely(EXCEL_REQUIREMENTS)
            for c in ["ID", "Category", "Group", "Requirement", "Explanation", "Verification", "Validation", "Related Parts"]:
                if c not in df.columns:
                    df[c] = ""
                else:
                    df[c] = df[c].fillna("").astype(object)
            return df
        except PermissionError:
            st.error(f"Close '{os.path.basename(EXCEL_REQUIREMENTS)}' in Microsoft Excel to continue.")
            return pd.DataFrame(columns=[
                "ID", "Category", "Group", "Requirement", "Explanation", "Verification", "Validation", "Related Parts"
            ])
    else:
        try:
            from init_requirements import INITIAL_REQUIREMENTS
            df = pd.DataFrame(INITIAL_REQUIREMENTS)
        except Exception:
            df = pd.DataFrame(columns=[
                "ID", "Category", "Group", "Requirement", "Explanation", "Verification", "Validation", "Related Parts"
            ])
        save_requirements(df)
        return df

def save_requirements(df):
    return save_formatted_excel(df, EXCEL_REQUIREMENTS, table_name="RequirementsTable")

# ==================== STREAMLIT CONFIG & STATE ====================
st.set_page_config(page_title="Drivetrain Hub", layout="wide")

if 'current_page' not in st.session_state:
    st.session_state.current_page = "Wheel Assembly"
elif st.session_state.current_page in ["Knowledge Hub", "Design Hub", "Assembly Tree", "Bearing Sizer"]:
    st.session_state.current_page = "Wheel Assembly"

if 'action' not in st.session_state:
    st.session_state.action = "view"
if 'selected_part' not in st.session_state:
    st.session_state.selected_part = None

df_parts = load_parts()
df_assemblies = load_assemblies()
df_trash = load_trash()
df_requirements = load_requirements()

# Global styling
st.markdown("""
<style>
div[data-testid="stSidebar"] button {
    justify-content: flex-start !important;
}
div[data-testid="stFullScreenFrame"] img {
    max-width: 95vw !important;
    max-height: 95vh !important;
    width: auto !important;
    height: auto !important;
    object-fit: contain !important;
}
textarea {
    field-sizing: content !important;
    min-height: 90px !important;
    max-height: 700px !important;
    line-height: 1.5 !important;
    resize: vertical !important;
}
</style>
""", unsafe_allow_html=True)

# ==================== SIDEBAR ====================
st.sidebar.markdown("### Drivetrain Hub")

# Wheel Assembly button
asm_active = (st.session_state.current_page == "Wheel Assembly")
if st.sidebar.button("Wheel Assembly", use_container_width=True, type="primary" if asm_active else "secondary"):
    st.session_state.current_page = "Wheel Assembly"
    st.session_state.selected_part = None
    st.session_state.action = "view"
    st.rerun()

# Requirements button
req_active = (st.session_state.current_page == "Requirements")
if st.sidebar.button("Requirements", use_container_width=True, type="primary" if req_active else "secondary"):
    st.session_state.current_page = "Requirements"
    st.session_state.selected_part = None
    st.session_state.action = "view"
    st.rerun()

st.sidebar.divider()

if st.sidebar.button("+ New Part", use_container_width=True):
    st.session_state.current_page = "Part Management"
    st.session_state.action = "create"
    st.session_state.selected_part = None
    st.rerun()

# Subassemblies
st.sidebar.markdown("### Components")
sub_labels = {
    "UP": "Uprights (UP)",
    "HB": "Hubs (HB)",
    "GB": "Gearbox (GB)"
}
custom_subs = [s for s in ["UP", "HB", "GB"] if s in df_parts["Subassembly"].unique()] if not df_parts.empty else []

for sub in custom_subs:
    sub_title = sub_labels.get(sub, sub)
    with st.sidebar.expander(sub_title, expanded=True):
        parts_in_sub = df_parts[df_parts["Subassembly"] == sub]["Part Name"].unique()
        for p in parts_in_sub:
            is_active = (st.session_state.current_page == "Part Management" and st.session_state.selected_part == p)
            if st.button(str(p), key=f"btn_{p}", use_container_width=True, type="primary" if is_active else "secondary"):
                st.session_state.current_page = "Part Management"
                st.session_state.selected_part = p
                st.session_state.action = "view"
                st.rerun()

# OEM Catalog
if not df_parts.empty and "Other" in df_parts["Subassembly"].unique():
    with st.sidebar.expander("OEM Parts", expanded=True):
        oem_parts = df_parts[df_parts["Subassembly"] == "Other"]["Part Name"].unique()
        for p in oem_parts:
            is_active = (st.session_state.current_page == "Part Management" and st.session_state.selected_part == p)
            if st.button(str(p), key=f"btn_oem_{p}", use_container_width=True, type="primary" if is_active else "secondary"):
                st.session_state.current_page = "Part Management"
                st.session_state.selected_part = p
                st.session_state.action = "view"
                st.rerun()

# Trash
st.sidebar.divider()
with st.sidebar.expander("Recycle Bin", expanded=False):
    if not df_trash.empty:
        trash_parts = df_trash["Part Name"].unique()
        for p in trash_parts:
            col1, col2 = st.columns([3, 2])
            col1.write(p)
            if col2.button("Restore", key=f"restore_{p}"):
                restored_rows = df_trash[df_trash["Part Name"] == p]
                df_parts = pd.concat([df_parts, restored_rows], ignore_index=True)
                save_parts(df_parts)
                df_trash = df_trash[df_trash["Part Name"] != p]
                save_trash(df_trash)
                st.rerun()
                
        if st.button("Empty Trash", use_container_width=True, type="primary"):
            df_trash = df_trash.iloc[0:0]
            save_trash(df_trash)
            st.rerun()
    else:
        st.caption("Trash is empty.")


# ==================== REQUIREMENTS COMPONENT ====================
def render_requirements_section(df_req):
    if df_req.empty:
        st.info("No requirements recorded.")
        return

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Rules", len(df_req))
    c2.metric("Team Rules", len(df_req[df_req["Category"] == "Team"]))
    c3.metric("FSG Rules", len(df_req[df_req["Category"] == "FSG"]))
    c4.metric("Inactive Rules", len(df_req[df_req["Category"] == "Inactive"]))

    st.divider()

    col_s1, col_s2, col_s3 = st.columns([2, 1, 1])
    with col_s1:
        search_kw = st.text_input("Search", placeholder="Search ID or text...", key="req_search")
    with col_s2:
        cat_filter = st.selectbox("Category", ["All", "Team", "FSG", "Inactive"], key="req_cat_filter")
    with col_s3:
        all_verif = sorted([str(x) for x in df_req["Verification"].dropna().unique() if str(x).strip() and str(x) != "-"])
        verif_filter = st.selectbox("Verification", ["All"] + all_verif, key="req_verif_filter")

    filtered_df = df_req.copy()
    if cat_filter != "All":
        filtered_df = filtered_df[filtered_df["Category"] == cat_filter]
    if verif_filter != "All":
        filtered_df = filtered_df[filtered_df["Verification"].astype(str) == verif_filter]
    if search_kw:
        kw = search_kw.lower().strip()
        filtered_df = filtered_df[
            filtered_df["ID"].astype(str).str.lower().str.contains(kw) |
            filtered_df["Requirement"].astype(str).str.lower().str.contains(kw) |
            filtered_df["Explanation"].astype(str).str.lower().str.contains(kw) |
            filtered_df["Related Parts"].astype(str).str.lower().str.contains(kw) |
            filtered_df["Group"].astype(str).str.lower().str.contains(kw)
        ]

    req_tabs = st.tabs([
        "All Rules",
        "Team Rules",
        "FSG Rules",
        "Inactive Rules",
        "Source Sheets"
    ])

    cols_to_show = ["ID", "Category", "Group", "Requirement", "Explanation", "Verification", "Validation", "Related Parts"]

    with req_tabs[0]:
        st.dataframe(filtered_df[cols_to_show], hide_index=True, use_container_width=True)

    with req_tabs[1]:
        team_df = filtered_df[filtered_df["Category"] == "Team"] if cat_filter == "All" else filtered_df
        st.dataframe(team_df[cols_to_show], hide_index=True, use_container_width=True)

    with req_tabs[2]:
        fsg_df = filtered_df[filtered_df["Category"] == "FSG"] if cat_filter == "All" else filtered_df
        st.dataframe(fsg_df[cols_to_show], hide_index=True, use_container_width=True)

    with req_tabs[3]:
        inact_df = filtered_df[filtered_df["Category"] == "Inactive"] if cat_filter == "All" else filtered_df
        st.dataframe(inact_df[cols_to_show], hide_index=True, use_container_width=True)

    with req_tabs[4]:
        st.markdown("#### Team Rules")
        col_img1, col_img2 = st.columns(2)
        with col_img1:
            p1 = os.path.join(IMAGE_DIR, "requirements", "team_requirements_p1.png")
            if os.path.exists(p1):
                st.image(p1, caption="Team Targets", use_container_width=True)
            p3 = os.path.join(IMAGE_DIR, "requirements", "team_requirements_p3.png")
            if os.path.exists(p3):
                st.image(p3, caption="Performance & Manufacturing", use_container_width=True)
        with col_img2:
            p2 = os.path.join(IMAGE_DIR, "requirements", "team_requirements_p2.png")
            if os.path.exists(p2):
                st.image(p2, caption="Assembly & Structural", use_container_width=True)

        st.markdown("#### FSG Rules")
        col_img3, col_img4 = st.columns(2)
        with col_img3:
            p4 = os.path.join(IMAGE_DIR, "requirements", "fsg_requirements_p1.png")
            if os.path.exists(p4):
                st.image(p4, caption="FSG General", use_container_width=True)
        with col_img4:
            p5 = os.path.join(IMAGE_DIR, "requirements", "fsg_requirements_p2.png")
            if os.path.exists(p5):
                st.image(p5, caption="FSG Fasteners & Inactive", use_container_width=True)

    st.divider()

    # Component options for Related Parts dropdown
    component_options = sorted(list(set(
        [str(p).strip() for p in df_parts["Part Name"].dropna().unique() if str(p).strip()] +
        [item.strip() for x in df_req["Related Parts"].dropna() for item in str(x).split(",") if item.strip()]
    )))

    col_act1, col_act2 = st.columns(2)
    with col_act1:
        with st.expander("Add Requirement"):
            with st.form("form_add_req"):
                new_id = st.text_input("ID")
                new_cat = st.selectbox("Category", ["Team", "FSG", "Inactive"])
                new_grp = st.text_input("Group")
                new_req = st.text_area("Requirement")
                new_exp = st.text_area("Explanation")
                new_ver = st.text_input("Verification")
                new_val = st.text_input("Validation")
                new_rel_parts = st.multiselect("Related Parts", options=component_options, default=[], key="add_req_rel_parts")
                if st.form_submit_button("Save", type="primary"):
                    if new_id.strip() and new_req.strip():
                        new_row = {
                            "ID": new_id.strip(),
                            "Category": new_cat,
                            "Group": new_grp.strip(),
                            "Requirement": new_req.strip(),
                            "Explanation": new_exp.strip(),
                            "Verification": new_ver.strip(),
                            "Validation": new_val.strip(),
                            "Related Parts": ", ".join(new_rel_parts)
                        }
                        updated_df = pd.concat([df_req, pd.DataFrame([new_row])], ignore_index=True)
                        save_requirements(updated_df)
                        st.success("Requirement added.")
                        st.rerun()
                    else:
                        st.error("Fields required.")

    with col_act2:
        with st.expander("Edit Requirement"):
            rule_ids = df_req["ID"].tolist()
            if rule_ids:
                selected_rule_id = st.selectbox("Select Rule", rule_ids)
                rule_row = df_req[df_req["ID"] == selected_rule_id].iloc[0]
                with st.form(f"form_edit_req_{selected_rule_id}"):
                    edit_cat = st.selectbox("Category", ["Team", "FSG", "Inactive"], index=["Team", "FSG", "Inactive"].index(rule_row["Category"]) if rule_row["Category"] in ["Team", "FSG", "Inactive"] else 0)
                    edit_grp = st.text_input("Group", value=str(rule_row.get("Group", "")))
                    edit_req = st.text_area("Requirement", value=str(rule_row.get("Requirement", "")))
                    edit_exp = st.text_area("Explanation", value=str(rule_row.get("Explanation", "")))
                    edit_ver = st.text_input("Verification", value=str(rule_row.get("Verification", "")))
                    edit_val = st.text_input("Validation", value=str(rule_row.get("Validation", "")))
                    current_rel_str = str(rule_row.get("Related Parts", ""))
                    current_rel_list = [p.strip() for p in current_rel_str.split(",") if p.strip()]
                    edit_options = sorted(list(set(component_options + current_rel_list)))
                    edit_rel_parts = st.multiselect("Related Parts", options=edit_options, default=current_rel_list, key=f"edit_req_rel_parts_{selected_rule_id}")
                    if st.form_submit_button("Save", type="primary"):
                        idx = df_req[df_req["ID"] == selected_rule_id].index[0]
                        df_req.at[idx, "Category"] = edit_cat
                        df_req.at[idx, "Group"] = edit_grp.strip()
                        df_req.at[idx, "Requirement"] = edit_req.strip()
                        df_req.at[idx, "Explanation"] = edit_exp.strip()
                        df_req.at[idx, "Verification"] = edit_ver.strip()
                        df_req.at[idx, "Validation"] = edit_val.strip()
                        df_req.at[idx, "Related Parts"] = ", ".join(edit_rel_parts)
                        save_requirements(df_req)
                        st.success("Requirement updated.")
                        st.rerun()

                if st.button("Delete Rule", key=f"del_rule_{selected_rule_id}"):
                    df_req_new = df_req[df_req["ID"] != selected_rule_id]
                    save_requirements(df_req_new)
                    st.success("Requirement deleted.")
                    st.rerun()


# =========================================================================
# PAGE 1: WHEEL ASSEMBLY
# =========================================================================
if st.session_state.current_page == "Wheel Assembly":
    st.title("Wheel Assembly")
    
    selected_parts_df = pd.DataFrame()
    ref_parts_df = pd.DataFrame()
    
    if not df_assemblies.empty:
        col1, col2 = st.columns(2)
        with col1:
            asm_versions_ref = df_assemblies["Assembly Version"].tolist()
            default_ref_idx = asm_versions_ref.index("FST15") if "FST15" in asm_versions_ref else 0
            ref_asm_v = st.selectbox("Baseline Assembly", asm_versions_ref, index=default_ref_idx)
        with col2:
            asm_versions_curr = df_assemblies["Assembly Version"].tolist()
            default_curr_idx = len(asm_versions_curr) - 1 if len(asm_versions_curr) > 1 else 0
            if default_curr_idx == default_ref_idx and len(asm_versions_curr) > 1:
                default_curr_idx = 0
            selected_asm_v = st.selectbox("Evaluate Assembly", asm_versions_curr, index=default_curr_idx)
            
        ref_row = df_assemblies[df_assemblies["Assembly Version"] == ref_asm_v].iloc[0]
        ref_config = parse_config(ref_row["Part Configuration"])
        for p_name, p_ver in ref_config.items():
            if p_name in ["Inner Bearing", "Outer Bearing", "Seal"]:
                actual_name = p_ver
                actual_ver = "OEM"
            else:
                actual_name = p_name
                actual_ver = p_ver
                
            row = df_parts[(df_parts["Part Name"].astype(str) == str(actual_name)) & (df_parts["Version"].astype(str) == str(actual_ver))]
            if not row.empty:
                row_copy = row.copy()
                row_copy["Part Name"] = p_name
                row_copy["Version"] = actual_name
                ref_parts_df = pd.concat([ref_parts_df, row_copy.iloc[[-1]]])
        
        asm_row = df_assemblies[df_assemblies["Assembly Version"] == selected_asm_v].iloc[0]
        if str(asm_row['Description']).strip() and str(asm_row['Description']) != "nan":
            st.info(asm_row['Description'])
            
        config = parse_config(asm_row["Part Configuration"])
        for p_name, p_ver in config.items():
            if p_name in ["Inner Bearing", "Outer Bearing", "Seal"]:
                actual_name = p_ver
                actual_ver = "OEM"
            else:
                actual_name = p_name
                actual_ver = p_ver
                
            row = df_parts[(df_parts["Part Name"].astype(str) == str(actual_name)) & (df_parts["Version"].astype(str) == str(actual_ver))]
            if not row.empty:
                row_copy = row.copy()
                row_copy["Part Name"] = p_name
                row_copy["Version"] = actual_name
                selected_parts_df = pd.concat([selected_parts_df, row_copy.iloc[[-1]]])
    else:
        st.info("No assemblies recorded.")
        if not df_parts.empty:
            for p in df_parts["Part Name"].unique():
                part_latest = df_parts[df_parts["Part Name"] == p].sort_values("Date").iloc[[-1]]
                selected_parts_df = pd.concat([selected_parts_df, part_latest])
                
        for p in df_parts["Part Name"].unique():
            part_first = df_parts[df_parts["Part Name"] == p].sort_values("Date").iloc[[0]]
            ref_parts_df = pd.concat([ref_parts_df, part_first])
            ref_asm_v = "Baseline"

    if not df_parts.empty:
        dashboard_data = []
        all_displayed_parts = set()
        if not ref_parts_df.empty: all_displayed_parts.update(ref_parts_df["Part Name"].astype(str).tolist())
        if not selected_parts_df.empty: all_displayed_parts.update(selected_parts_df["Part Name"].astype(str).tolist())
        
        col_in_ref = f"In {ref_asm_v}"
        col_in_curr = f"In {selected_asm_v}"
        if col_in_ref == col_in_curr:
            col_in_ref = f"In {ref_asm_v} (Ref)"
            col_in_curr = f"In {selected_asm_v} (Eval)"

        for p_name in sorted(list(all_displayed_parts)):
            if not ref_parts_df.empty and str(p_name) in ref_parts_df["Part Name"].astype(str).values:
                orig_weight = ref_parts_df[ref_parts_df["Part Name"].astype(str) == str(p_name)].iloc[0]["Weight (g)"]
            else:
                orig_weight = 0.0
                
            if not selected_parts_df.empty and str(p_name) in selected_parts_df["Part Name"].astype(str).values:
                curr_row = selected_parts_df[selected_parts_df["Part Name"].astype(str) == str(p_name)].iloc[0]
                curr_weight = curr_row["Weight (g)"]
                subassembly = curr_row["Subassembly"]
            else:
                curr_weight = 0.0
                if not ref_parts_df.empty and str(p_name) in ref_parts_df["Part Name"].astype(str).values:
                    subassembly = ref_parts_df[ref_parts_df["Part Name"].astype(str) == str(p_name)].iloc[0]["Subassembly"]
                else:
                    subassembly = "Other"
                
            if orig_weight == 0.0 and curr_weight == 0.0:
                continue
                
            diff = curr_weight - orig_weight
            in_ref = "Yes" if (not ref_parts_df.empty and str(p_name) in ref_parts_df["Part Name"].astype(str).values) else "No"
            in_curr = "Yes" if (not selected_parts_df.empty and str(p_name) in selected_parts_df["Part Name"].astype(str).values) else "No"
            
            dashboard_data.append({
                "Subassembly": subassembly,
                "Part Name": p_name,
                col_in_ref: in_ref,
                col_in_curr: in_curr,
                "Ref. Weight (g)": orig_weight,
                "Current Weight (g)": curr_weight,
                "Diff (g)": diff
            })
            
        dash_df = pd.DataFrame(dashboard_data, columns=[
            "Subassembly", "Part Name", 
            col_in_ref, col_in_curr, 
            "Ref. Weight (g)", "Current Weight (g)", "Diff (g)"
        ]).reset_index(drop=True)
        
        total_curr = dash_df["Current Weight (g)"].sum() if not dash_df.empty else 0.0
        total_orig = dash_df["Ref. Weight (g)"].sum() if not dash_df.empty else 0.0
        total_diff = total_curr - total_orig
        
        col1, col2 = st.columns(2)
        col1.metric("Reference Weight", f"{total_orig:.2f} g")
        col2.metric("Current Weight", f"{total_curr:.2f} g", delta=f"{total_diff:+.2f} g", delta_color="inverse")
        
        st.divider()
        st.subheader("Part Comparison")
        
        def highlight_diff(val):
            try:
                v = float(val)
                if v < -0.0001:
                    return 'color: #33cc33; font-weight: 600;'
                elif v > 0.0001:
                    return 'color: #ff4b4b; font-weight: 600;'
            except Exception:
                pass
            return ''
            
        styled_df = dash_df.style.map(highlight_diff, subset=['Diff (g)']).format({
            "Current Weight (g)": "{:.2f}",
            "Ref. Weight (g)": "{:.2f}",
            "Diff (g)": "{:+.2f}"
        })
        
        st.dataframe(styled_df, use_container_width=True, hide_index=True)

        if not selected_parts_df.empty:
            with st.expander("Component Summary", expanded=False):
                active_part_rows = []
                for _, s_row in selected_parts_df.iterrows():
                    role_name = str(s_row["Part Name"])
                    ver_str = str(s_row["Version"])
                    real_name = ver_str if role_name in ["Inner Bearing", "Outer Bearing", "Seal"] else role_name
                    active_part_rows.append((role_name, real_name, ver_str, s_row))

                summary_items = []
                for role_name, real_name, ver_str, s_row in active_part_rows:
                    just_txt = str(s_row.get("Justification", "")).strip()
                    if just_txt == "nan":
                        just_txt = ""
                    tols_list = parse_tolerancing_data(s_row.get("Tolerancing Data", ""))
                    reqs_df = pd.DataFrame()
                    if not df_requirements.empty:
                        m_req = df_requirements["Related Parts"].astype(str).str.contains(str(real_name), na=False, regex=False)
                        reqs_df = df_requirements[m_req]
                    if just_txt or tols_list or not reqs_df.empty:
                        summary_items.append((role_name, real_name, ver_str, just_txt, tols_list, reqs_df))

                if summary_items:
                    comp_names = [f"{r_name}" if r_name == real_n else f"{r_name} ({real_n})" for r_name, real_n, _, _, _, _ in summary_items]
                    sel_comp_idx = st.selectbox("Select Component", range(len(comp_names)), format_func=lambda i: comp_names[i], key="asm_comp_summary_sel")
                    _, real_n, _, j_txt, t_list, r_df = summary_items[sel_comp_idx]
                    if j_txt:
                        st.markdown("##### Justification")
                        st.markdown(j_txt)
                    if t_list:
                        st.markdown("##### Tolerances")
                        st.dataframe(pd.DataFrame(t_list), hide_index=True, use_container_width=True)
                    if not r_df.empty:
                        st.markdown("##### Related Requirements")
                        st.dataframe(r_df[["ID", "Category", "Requirement", "Verification", "Validation"]], hide_index=True, use_container_width=True)
                else:
                    st.caption("No component notes recorded.")

    st.divider()
    with st.expander("Manage Assemblies"):
        mode = st.radio("Action", ["Create Version", "Edit Assembly"], horizontal=True)
        is_edit = mode == "Edit Assembly"
        
        target_asm = None
        if is_edit:
            if df_assemblies.empty:
                st.warning("No assemblies to edit.")
            else:
                target_asm = st.selectbox("Select Assembly", df_assemblies["Assembly Version"].tolist())
                
        if not is_edit or (is_edit and not df_assemblies.empty and target_asm):
            with st.form("new_asm_form"):
                if is_edit:
                    target_row = df_assemblies[df_assemblies["Assembly Version"] == target_asm].iloc[0]
                    new_asm_v = st.text_input("Version ID", value=target_asm, disabled=True)
                    new_asm_desc = st.text_area("Description", value=target_row["Description"], height=calc_description_height(target_row["Description"]))
                    base_config = parse_config(target_row["Part Configuration"])
                else:
                    new_asm_v = st.text_input("Version ID", value="1.0" if df_assemblies.empty else "")
                    new_asm_desc = st.text_area("Description", height=calc_description_height(""))
                    base_config = {}
                    if not df_assemblies.empty:
                        base_config = parse_config(df_assemblies.iloc[-1]["Part Configuration"])
                
                st.markdown("#### Part Configuration")
                
                new_config = {}
                unique_parts = df_parts["Part Name"].unique().tolist() if not df_parts.empty else []
                
                bearings_list = []
                seals_list = []
                custom_parts = []
                
                if not df_parts.empty:
                    oem_mask = df_parts["Subassembly"].astype(str) == "Other"
                    oem_parts_names = df_parts[oem_mask]["Part Name"].unique().tolist()
                    
                    for p in oem_parts_names:
                        p_str = str(p).upper()
                        row_match = df_parts[df_parts["Part Name"].astype(str) == str(p)]
                        desc_str = str(row_match.iloc[0].get("Description", "")).upper() if not row_match.empty else ""
                        
                        if "HMSA" in p_str or "SEAL" in p_str or "SKF" in desc_str or "SEAL" in desc_str or ("X" in p_str and any(c.isdigit() for c in p_str)):
                            seals_list.append(str(p))
                        else:
                            bearings_list.append(str(p))
                            
                    for p in unique_parts:
                        if str(p) not in oem_parts_names:
                            custom_parts.append(p)
                
                if len(custom_parts) > 0:
                    st.markdown("##### Custom Components")
                    for p in custom_parts:
                        raw_versions = df_parts[df_parts["Part Name"].astype(str) == str(p)]["Version"].tolist()
                        p_versions = ["-- Not Used --"] + raw_versions
                        
                        default_v = base_config.get(p, "-- Not Used --")
                        if default_v not in p_versions:
                            default_v = raw_versions[-1]
                            
                        idx = p_versions.index(default_v)
                        sel_v = st.selectbox(f"{p}", p_versions, index=idx)
                        if sel_v != "-- Not Used --":
                            new_config[p] = sel_v
                else:
                    st.warning("No custom parts exist.")
                    
                st.markdown("##### OEM Components")
                b_opts = ["-- Not Used --"] + [str(b) for b in bearings_list]
                s_opts = ["-- Not Used --"] + [str(s) for s in seals_list]
                
                existing_inner = base_config.get("Inner Bearing", "-- Not Used --")
                existing_outer = base_config.get("Outer Bearing", "-- Not Used --")
                existing_seal = base_config.get("Seal", "-- Not Used --")
                
                def match_oem_opt(val, options):
                    if not val or val == "-- Not Used --":
                        return 0
                    if val in options:
                        return options.index(val)
                    for idx_o, opt_val in enumerate(options):
                        if opt_val != "-- Not Used --":
                            if opt_val in val or val in opt_val:
                                return idx_o
                    return 0

                idx_in = match_oem_opt(existing_inner, b_opts)
                idx_out = match_oem_opt(existing_outer, b_opts)
                idx_s = match_oem_opt(existing_seal, s_opts)
                
                sel_inner = st.selectbox("Inner Bearing", b_opts, index=idx_in)
                sel_outer = st.selectbox("Outer Bearing", b_opts, index=idx_out)
                sel_seal = st.selectbox("Seal", s_opts, index=idx_s)
                
                if sel_inner != "-- Not Used --": new_config["Inner Bearing"] = sel_inner
                if sel_outer != "-- Not Used --": new_config["Outer Bearing"] = sel_outer
                if sel_seal != "-- Not Used --": new_config["Seal"] = sel_seal
                
                submitted = st.form_submit_button("Save Assembly", type="primary")
                if submitted:
                    if not new_asm_v.strip():
                        st.error("Version ID required.")
                    else:
                        total_weight = 0.0
                        for p_name, p_ver in new_config.items():
                            if p_name in ["Inner Bearing", "Outer Bearing", "Seal"]:
                                actual_name = p_ver
                                actual_ver = "OEM"
                            else:
                                actual_name = p_name
                                actual_ver = p_ver
                                
                            part_match = df_parts[(df_parts["Part Name"].astype(str) == str(actual_name)) & (df_parts["Version"].astype(str) == str(actual_ver))]
                            if not part_match.empty:
                                total_weight += float(part_match.iloc[-1]["Weight (g)"])
                                
                        new_row = {
                            "Assembly Version": new_asm_v,
                            "Date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                            "Total Weight (g)": round(total_weight, 2),
                            "Description": new_asm_desc,
                            "Part Configuration": format_config(new_config)
                        }
                        
                        if is_edit:
                            df_assemblies = df_assemblies[df_assemblies["Assembly Version"] != target_asm]
                            
                        df_assemblies = pd.concat([df_assemblies, pd.DataFrame([new_row])], ignore_index=True)
                        save_assemblies(df_assemblies)
                        st.success("Assembly saved successfully.")
                        st.rerun()




# =========================================================================
# PAGE 2: REQUIREMENTS
# =========================================================================
elif st.session_state.current_page == "Requirements":
    st.title("Requirements")
    render_requirements_section(df_requirements)


# =========================================================================
# PAGE 3: PART MANAGEMENT
# =========================================================================
elif st.session_state.current_page == "Part Management":
    
    # ------------------ CREATE PART ------------------
    if st.session_state.action == "create":
        st.title("New Part")
        
        with st.form("new_part_form"):
            col1, col2 = st.columns(2)
            with col1:
                sub_opts_map = {
                    "Uprights (UP)": "UP",
                    "Hubs (HB)": "HB",
                    "Gearbox (GB)": "GB",
                    "OEM Parts": "Other"
                }
                sub_choice = st.selectbox("Subassembly", list(sub_opts_map.keys()))
                subassembly = sub_opts_map[sub_choice]
                part_name = st.text_input("Part Name")
            with col2:
                initial_version = st.text_input("Initial Version", value="1.0" if subassembly != "Other" else "OEM")
                initial_weight = st.number_input("Weight (g)", min_value=0.0, format="%.2f")
            
            initial_desc = st.text_area("Description", height=calc_description_height(""))
            c_opt1, c_opt2 = st.columns(2)
            with c_opt1:
                has_just_create = st.checkbox("Include Justification", value=False, key="cb_create_has_just")
            with c_opt2:
                has_tol_create = st.checkbox("Include Tolerances", value=False, key="cb_create_has_tol")
            if has_just_create:
                create_just = st.text_area("Justification", height=100, key="create_just_text")
            else:
                create_just = ""
            
            st.divider()
            st.subheader("Specifications")
            col_tables, col_img = st.columns([1.5, 1.2])
            with col_tables:
                st.markdown("**Specifications**")
                meas_df_init = pd.DataFrame([{"Property": "", "Value": ""}])
                edited_meas_init = st.data_editor(meas_df_init, num_rows="dynamic", use_container_width=True, key="new_part_meas")

                if has_tol_create:
                    st.markdown("**Tolerances**")
                    tol_df_init = pd.DataFrame([{"Interface / Feature": "", "Nominal (mm)": "", "Fit / Tolerance": "", "Roughness / Notes": ""}])
                    edited_tol_init = st.data_editor(tol_df_init, num_rows="dynamic", use_container_width=True, key="new_part_tol")
                else:
                    edited_tol_init = pd.DataFrame()
                
            with col_img:
                pasted_img = render_image_input("Schematic", key="new_part_img")

            submitted = st.form_submit_button("Create Part", type="primary")
            if submitted:
                final_part_name = part_name.strip()
                if final_part_name:
                    initial_version = str(initial_version).lstrip('vV')
                    img_path = save_pasted_image(pasted_img, final_part_name, initial_version)
                    if img_path:
                        st.session_state.pop("clipboard_new_part_img", None)
                        
                    final_meas = [r for r in edited_meas_init.to_dict('records') if str(r.get('Property', '')).strip() != ""]
                    final_tol = [r for r in edited_tol_init.to_dict('records') if str(r.get('Interface / Feature', '')).strip() != ""] if (has_tol_create and not edited_tol_init.empty) else []

                    new_row = {
                        "Subassembly": subassembly,
                        "Part Name": final_part_name,
                        "Version": initial_version,
                        "Date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "Weight (g)": initial_weight,
                        "Description": initial_desc,
                        "Specifications": format_measurements(final_meas),
                        "Tolerancing Data": format_tolerancing_data(final_tol),
                        "Image Path": img_path,
                        "Stiffness Data": "",
                        "LC Images": "",
                        "Justification": create_just.strip() if has_just_create else ""
                    }
                    df_parts = pd.concat([df_parts, pd.DataFrame([new_row])], ignore_index=True)
                    save_parts(df_parts)
                    st.session_state.action = "view"
                    st.session_state.selected_part = final_part_name
                    st.success(f"Component '{final_part_name}' created successfully.")
                    st.rerun()
                else:
                    st.error("Part Name required.")

    # ------------------ VIEW PART ------------------
    elif st.session_state.action == "view" and st.session_state.selected_part:
        selected_part = st.session_state.selected_part
        
        if df_parts.empty or selected_part not in df_parts["Part Name"].values:
            st.warning("Part not found.")
            if st.button("Wheel Assembly"):
                st.session_state.current_page = "Wheel Assembly"
                st.session_state.selected_part = None
                st.rerun()
        else:
            part_data = df_parts[df_parts["Part Name"] == selected_part].sort_values("Date")
            latest_iter = part_data.iloc[-1]
            subassembly = latest_iter['Subassembly']
            is_oem = (subassembly == "Other") or (str(selected_part) in ["61815", "61816", "71815-TN", "71816-TN", "85X100X9 HMSA10 V", "95X110X9 HMSA10 V"])
            is_structural = (str(selected_part) in ["UP_01", "UP_02", "UpperSupport Version", "HB_01"])

            if f"pending_eval_{selected_part}" in st.session_state:
                st.session_state[f"comp_eval_{selected_part}"] = st.session_state.pop(f"pending_eval_{selected_part}")
            if f"pending_edit_{selected_part}" in st.session_state:
                st.session_state[f"sel_edit_{selected_part}"] = st.session_state.pop(f"pending_edit_{selected_part}")

            st.title(str(selected_part))

            # Metrics
            if is_oem:
                col1, col2, col3 = st.columns(3)
                col1.metric("Version", f"{latest_iter['Version']}")
                col2.metric("Weight", f"{latest_iter['Weight (g)']} g")
                col3.metric("Subassembly", "OEM Parts")
            else:
                col1, col2, col3, col4 = st.columns(4)
                col1.metric("Version", f"{format_v_label(latest_iter['Version'])}")
                col2.metric("Weight", f"{latest_iter['Weight (g)']} g")
                weight_diff = latest_iter['Weight (g)'] - part_data.iloc[0]['Weight (g)']
                col3.metric("Net Diff", f"{weight_diff:+.2f} g")
                sub_disp = {"UP": "Uprights (UP)", "HB": "Hubs (HB)", "GB": "Gearbox (GB)"}.get(subassembly, subassembly)
                col4.metric("Subassembly", sub_disp)

            # Clean Justification Note at Top (only rendered if justification exists)
            curr_just = str(latest_iter.get("Justification", "")).strip()
            if curr_just == "nan":
                curr_just = ""

            if curr_just:
                if "\n" not in curr_just and len(curr_just) <= 180:
                    st.caption(f"*{curr_just}*")
                else:
                    first_line = curr_just.strip().split("\n")[0].lstrip("#").strip()
                    if first_line:
                        st.caption(f"*{first_line}*")

            # Clean Tolerances Note at Top (only rendered if tolerances exist)
            curr_tols = parse_tolerancing_data(latest_iter.get("Tolerancing Data", ""))
            if curr_tols:
                tol_features = ", ".join([str(t.get('Interface / Feature', '')).strip() for t in curr_tols if str(t.get('Interface / Feature', '')).strip()])
                if tol_features:
                    if len(tol_features) > 100:
                        tol_features = tol_features[:97] + "..."
                    st.caption(f"*Tolerances: {tol_features}*")

            # Tabs: Overview, Justification, Tolerances, New Iteration, Edit
            tab_overview, tab_just, tab_tol, tab_nova, tab_editar = st.tabs([
                "Overview", 
                "Justification", 
                "Tolerances",
                "New Iteration", 
                "Edit"
            ])

            # ---------- TAB: OVERVIEW ----------
            with tab_overview:
                if not is_oem and len(part_data) > 1:
                    st.subheader("Version Comparison")
                    comp_versions = []
                    for v in part_data["Version"].astype(str).tolist():
                        if v not in comp_versions:
                            comp_versions.append(v)

                    col1, col2 = st.columns(2)
                    ref_key = f"comp_ref_{selected_part}"
                    if ref_key in st.session_state and st.session_state[ref_key] not in comp_versions:
                        st.session_state.pop(ref_key, None)
                    eval_key = f"comp_eval_{selected_part}"
                    if eval_key in st.session_state and st.session_state[eval_key] not in comp_versions:
                        st.session_state.pop(eval_key, None)

                    with col1:
                        default_ref_idx = comp_versions.index("FST15") if "FST15" in comp_versions else 0
                        ref_v = st.selectbox(
                            "Baseline Iteration",
                            comp_versions,
                            index=default_ref_idx,
                            format_func=format_v_label,
                            key=ref_key
                        )
                    with col2:
                        default_eval_idx = len(comp_versions) - 1 if len(comp_versions) > 1 else 0
                        eval_v = st.selectbox(
                            "Evaluate Iteration",
                            comp_versions,
                            index=default_eval_idx,
                            format_func=format_v_label,
                            key=eval_key
                        )

                    ref_rows = part_data[part_data["Version"].astype(str) == str(ref_v)]
                    ref_row = ref_rows.iloc[-1] if not ref_rows.empty else latest_iter
                    eval_rows = part_data[part_data["Version"].astype(str) == str(eval_v)]
                    eval_row = eval_rows.iloc[-1] if not eval_rows.empty else latest_iter

                    ref_label = format_v_label(ref_v)
                    eval_label = format_v_label(eval_v)

                    st.markdown("#### Mass Difference")
                    df_mass = build_mass_comparison_dataframe(
                        eval_weight=eval_row['Weight (g)'],
                        ref_weight=ref_row['Weight (g)'],
                        curr_label=eval_label,
                        ref_label=ref_label
                    )
                    st.dataframe(style_mass_dataframe(df_mass), use_container_width=True, hide_index=True)

                    if len(comp_versions) > 2:
                        with st.expander("All Iterations", expanded=False):
                            df_all_mass = build_all_iterations_mass_dataframe(part_data, ref_row['Weight (g)'], ref_label)
                            st.dataframe(style_mass_dataframe(df_all_mass), use_container_width=True, hide_index=True)

                    # Structural FEA Comparisons (modular for UP_01, HB_01 and brackets)
                    part_fea_cfg = get_part_fea_config(selected_part)
                    part_lcs = part_fea_cfg.get("LC_KEYS", [])
                    if part_lcs:
                        st.markdown("#### FEA Displacements")
                        part_info = part_fea_cfg.get("LC_INFO", {})
                        ref_stiff = parse_stiffness_data(ref_row.get("Stiffness Data", ""), lc_keys=part_lcs)
                        if not ref_stiff and str(ref_v) == "FST15":
                            ref_stiff = part_fea_cfg.get("DEFAULT_FST15", {})

                        eval_stiff = parse_stiffness_data(eval_row.get("Stiffness Data", ""), lc_keys=part_lcs)
                        if not eval_stiff:
                            if str(eval_v) == "FST15":
                                eval_stiff = part_fea_cfg.get("DEFAULT_FST15", {})
                            else:
                                eval_stiff = part_fea_cfg.get("DEFAULT_FST16", {})

                        df_view = build_comparison_dataframe(eval_stiff, ref_stiff, ref_label=ref_label, curr_label=eval_label, lc_info=part_info, lc_keys=part_lcs)
                        st.dataframe(style_comparison_dataframe(df_view, decimals=part_fea_cfg.get("DECIMALS", 4)), use_container_width=True, hide_index=True)

                        ref_lc_dict = parse_lc_images(ref_row.get("LC Images", ""))
                        eval_lc_dict = parse_lc_images(eval_row.get("LC Images", ""))
                        if eval_lc_dict or ref_lc_dict:
                            with st.expander("FEA Screenshots", expanded=False):
                                avail_lcs = [lc for lc in part_lcs if lc in eval_lc_dict or lc in ref_lc_dict]
                                if avail_lcs:
                                    view_tabs = st.tabs([f"{lc}" for lc in avail_lcs])
                                    for idx_lc, lc in enumerate(avail_lcs):
                                        with view_tabs[idx_lc]:
                                            desc_label = part_info.get(lc, {}).get("desc", lc)
                                            st.write(f"**{lc} - {desc_label}**")
                                            img_ref_f = resolve_image_path(ref_lc_dict.get(lc))
                                            img_eval_f = resolve_image_path(eval_lc_dict.get(lc))
                                            if str(ref_v) != str(eval_v) and img_ref_f and img_eval_f:
                                                c_lc1, c_lc2 = st.columns(2)
                                                with c_lc1:
                                                    st.image(img_ref_f, caption=f"{ref_label} - {lc}", use_container_width=True)
                                                with c_lc2:
                                                    st.image(img_eval_f, caption=f"{eval_label} - {lc}", use_container_width=True)
                                            elif img_eval_f:
                                                st.image(img_eval_f, caption=f"{eval_label} - {lc}", use_container_width=True)
                                            elif img_ref_f:
                                                st.image(img_ref_f, caption=f"{ref_label} - {lc}", use_container_width=True)
                    else:
                        eval_struct = parse_structural_data(eval_row.get("Stiffness Data", ""))
                        ref_struct = parse_structural_data(ref_row.get("Stiffness Data", ""))
                        if eval_struct or ref_struct:
                            st.markdown("#### Structural FEA")
                            e_sf = float(eval_struct.get("sf", 0.0))
                            r_sf = float(ref_struct.get("sf", 0.0))
                            diff_sf = e_sf - r_sf
                            e_stress = float(eval_struct.get("stress", 0.0))
                            r_stress = float(ref_struct.get("stress", 0.0))
                            diff_stress = e_stress - r_stress
                            e_disp = float(eval_struct.get("disp", 0.0))
                            r_disp = float(ref_struct.get("disp", 0.0))
                            diff_disp = e_disp - r_disp

                            col_b1, col_b2, col_b3, col_b4 = st.columns(4)
                            col_b1.metric("Safety Factor", f"{e_sf:.2f}" if e_sf > 0 else "-", delta=f"{diff_sf:+.2f} vs {ref_label}" if (e_sf > 0 and r_sf > 0) else None, delta_color="normal")
                            col_b2.metric("Peak Stress", f"{e_stress:.1f} MPa" if e_stress > 0 else "-", delta=f"{diff_stress:+.1f} MPa vs {ref_label}" if (e_stress > 0 and r_stress > 0) else None, delta_color="inverse")
                            col_b3.metric("Max Deflection", f"{e_disp:.4f} mm" if e_disp > 0 else "-", delta=f"{diff_disp:+.4f} mm vs {ref_label}" if (e_disp > 0 and r_disp > 0) else None, delta_color="inverse")
                            sf_status = "Meets Target" if e_sf >= 2.03 else ("Under Target" if e_sf > 0 else "Pending FEA")
                            col_b4.metric("Criteria", sf_status)

                            df_bracket_comp = build_bracket_comparison_dataframe(eval_struct, ref_struct, eval_label=eval_label, ref_label=ref_label)
                            st.dataframe(df_bracket_comp, hide_index=True, use_container_width=True)

                            eval_lc_dict = parse_lc_images(eval_row.get("LC Images", ""))
                            fea_img_key = next((k for k in ["FEA", "Structural", "LC1"] if k in eval_lc_dict), None)
                            if fea_img_key and eval_lc_dict.get(fea_img_key):
                                img_bracket = resolve_image_path(eval_lc_dict[fea_img_key])
                                if img_bracket:
                                    st.image(img_bracket, caption=f"{selected_part} ({eval_label})", use_container_width=True)

                    st.divider()
                elif not is_oem:
                    st.info("Single iteration recorded.")

                # Synced Side-by-Side Comparison when comparing two distinct iterations
                if not is_oem and len(part_data) > 1 and str(ref_v) != str(eval_v):
                    col_base, col_eval = st.columns(2)
                    for col_side, side_title, side_label, side_row in [
                        (col_base, f"Baseline: {ref_label}", ref_label, ref_row),
                        (col_eval, f"Evaluated: {eval_label}", eval_label, eval_row),
                    ]:
                        with col_side:
                            st.markdown(f"#### {side_title}")
                            st.caption(f"Weight: {float(side_row['Weight (g)']):.2f} g | Date: {side_row['Date']}")
                            side_desc = str(side_row.get("Description", "")).strip()
                            if side_desc and side_desc != "nan":
                                st.markdown(side_desc)
                            side_specs = parse_measurements(side_row.get("Specifications", side_row.get("Constraints", "")))
                            if side_specs:
                                st.markdown("**Specifications**")
                                st.dataframe(pd.DataFrame(side_specs), hide_index=True, use_container_width=True)
                            side_img = resolve_image_path(side_row.get("Image Path"))
                            if side_img:
                                st.image(side_img, caption=f"{selected_part} {side_label}", use_container_width=True)
                            else:
                                st.caption("No schematic image.")
                else:
                    single_row = eval_row if (not is_oem and len(part_data) > 1) else latest_iter
                    if not is_oem and len(part_data) > 1 and str(ref_v) == str(eval_v):
                        st.caption(f"Showing {eval_label}. Select distinct iterations above to compare.")
                    col_left, col_right = st.columns([1.5, 1.2])
                    with col_left:
                        st.markdown("#### Specifications")
                        single_desc = str(single_row.get("Description", "")).strip()
                        if single_desc and single_desc != "nan":
                            st.markdown(single_desc)
                        raw_specs = parse_measurements(single_row.get("Specifications", single_row.get("Constraints", "")))
                        if raw_specs:
                            st.dataframe(pd.DataFrame(raw_specs), hide_index=True, use_container_width=True)
                        else:
                            st.caption("No specifications recorded.")

                    with col_right:
                        st.markdown("#### Schematic")
                        img_f = resolve_image_path(single_row.get("Image Path"))
                        if img_f:
                            st.image(img_f, caption=f"{selected_part} Schematic", use_container_width=True)
                        else:
                            st.info("No schematic image.")

                if not df_requirements.empty and selected_part:
                    match_mask = df_requirements["Related Parts"].astype(str).str.contains(str(selected_part), na=False)
                    part_reqs = df_requirements[match_mask]
                    if not part_reqs.empty:
                        with st.expander("Related Requirements", expanded=False):
                            st.dataframe(
                                part_reqs[["ID", "Category", "Requirement", "Verification", "Validation"]],
                                hide_index=True,
                                use_container_width=True
                            )

            # ---------- TAB: JUSTIFICATION ----------
            with tab_just:
                st.subheader("Component Justification")
                if curr_just:
                    st.markdown(curr_just)
                    st.divider()

                with st.expander("Edit Justification", expanded=(not bool(curr_just))):
                    has_just = st.checkbox(
                        "Include Justification",
                        value=bool(curr_just),
                        key=f"tab_just_cb_{selected_part}"
                    )
                    if has_just:
                        new_just_text = st.text_area(
                            "Justification",
                            value=curr_just,
                            height=280,
                            key=f"tab_just_text_{selected_part}"
                        )
                    else:
                        new_just_text = ""

                    c_j1, c_j2 = st.columns(2)
                    with c_j1:
                        if st.button("Save", type="primary", use_container_width=True, key=f"btn_save_just_{selected_part}"):
                            if "Justification" in df_parts.columns and df_parts["Justification"].dtype != object:
                                df_parts["Justification"] = df_parts["Justification"].astype(object)
                            mask = df_parts["Part Name"] == selected_part
                            val_to_save = new_just_text.strip() if has_just else ""
                            df_parts.loc[mask, "Justification"] = val_to_save
                            save_parts(df_parts)
                            st.success("Justification saved.")
                            st.rerun()
                    with c_j2:
                        if curr_just:
                            if st.button("Remove", use_container_width=True, key=f"btn_rem_just_{selected_part}"):
                                if "Justification" in df_parts.columns and df_parts["Justification"].dtype != object:
                                    df_parts["Justification"] = df_parts["Justification"].astype(object)
                                mask = df_parts["Part Name"] == selected_part
                                df_parts.loc[mask, "Justification"] = ""
                                save_parts(df_parts)
                                st.success("Justification removed.")
                                st.rerun()

            # ---------- TAB: TOLERANCES ----------
            with tab_tol:
                st.subheader("Component Tolerances")
                curr_comp_tols = parse_tolerancing_data(latest_iter.get("Tolerancing Data", ""))
                has_tols = st.checkbox(
                    "Include Tolerances",
                    value=bool(curr_comp_tols),
                    key=f"tab_tol_cb_{selected_part}"
                )
                if has_tols:
                    tol_records = curr_comp_tols if curr_comp_tols else [{
                        "Interface / Feature": "",
                        "Nominal (mm)": "",
                        "Fit / Tolerance": "",
                        "Roughness / Notes": ""
                    }]
                    tol_df_edit = pd.DataFrame(tol_records)
                    edited_tols = st.data_editor(
                        tol_df_edit,
                        num_rows="dynamic",
                        use_container_width=True,
                        key=f"tab_tol_editor_{selected_part}"
                    )
                else:
                    st.caption("No tolerances defined for this component.")

                c_t1, c_t2 = st.columns(2)
                with c_t1:
                    if st.button("Save", type="primary", use_container_width=True, key=f"btn_save_tol_{selected_part}"):
                        final_tol = []
                        if has_tols:
                            for r in edited_tols.to_dict('records'):
                                if any(str(v).strip() for v in r.values()):
                                    feat = str(r.get("Interface / Feature", "")).strip()
                                    if feat:
                                        final_tol.append({
                                            "Interface / Feature": feat,
                                            "Nominal (mm)": str(r.get("Nominal (mm)", "")).strip(),
                                            "Fit / Tolerance": str(r.get("Fit / Tolerance", "")).strip(),
                                            "Roughness / Notes": str(r.get("Roughness / Notes", "")).strip()
                                        })
                        formatted_tol = format_tolerancing_data(final_tol) if final_tol else ""
                        if "Tolerancing Data" in df_parts.columns and df_parts["Tolerancing Data"].dtype != object:
                            df_parts["Tolerancing Data"] = df_parts["Tolerancing Data"].astype(object)
                        mask = df_parts["Part Name"] == selected_part
                        df_parts.loc[mask, "Tolerancing Data"] = formatted_tol
                        save_parts(df_parts)
                        st.success("Tolerances saved.")
                        st.rerun()
                with c_t2:
                    if curr_comp_tols:
                        if st.button("Remove", use_container_width=True, key=f"btn_rem_tol_{selected_part}"):
                            if "Tolerancing Data" in df_parts.columns and df_parts["Tolerancing Data"].dtype != object:
                                df_parts["Tolerancing Data"] = df_parts["Tolerancing Data"].astype(object)
                            mask = df_parts["Part Name"] == selected_part
                            df_parts.loc[mask, "Tolerancing Data"] = ""
                            save_parts(df_parts)
                            st.success("Tolerances removed.")
                            st.rerun()

            # ---------- TAB: NEW ITERATION ----------
            with tab_nova:
                st.subheader("New Iteration")
                source_indices = part_data.sort_values("Date", ascending=False).index.tolist()
                if len(source_indices) > 1:
                    def format_source_version(idx):
                        r = part_data.loc[idx]
                        return f"{format_v_label(r['Version'])} ({r['Date']})"
                    copy_key = f"copy_from_{selected_part}"
                    if copy_key in st.session_state and st.session_state[copy_key] not in source_indices:
                        st.session_state.pop(copy_key, None)
                    source_idx = st.selectbox(
                        "Copy From",
                        source_indices,
                        format_func=format_source_version,
                        index=0,
                        key=copy_key
                    )
                    source_iter = part_data.loc[source_idx]
                else:
                    source_idx = source_indices[0]
                    source_iter = latest_iter

                current_v = str(latest_iter['Version'])
                try:
                    if '.' in current_v:
                        parts_v = current_v.split('.')
                        major = int(parts_v[0])
                        minor = int(parts_v[-1])
                        prefix = ".".join(parts_v[:-1])
                        next_minor = f"{prefix}.{minor + 1}"
                    else:
                        major = int(current_v)
                        next_minor = f"{major}.1"
                except Exception:
                    next_minor = f"{current_v}_new"

                col1, col2 = st.columns(2)
                with col1:
                    new_version = st.text_input("Version ID", value=next_minor, key=f"new_ver_{selected_part}")
                    weight = st.number_input("Weight (g)", value=float(source_iter['Weight (g)']), format="%.2f", key=f"new_wt_{selected_part}_{source_idx}")
                with col2:
                    src_desc = str(source_iter.get('Description', ''))
                    if src_desc == "nan":
                        src_desc = ""
                    desc = st.text_area(
                        "Description",
                        value=src_desc,
                        height=calc_description_height(st.session_state.get(f"new_desc_{selected_part}_{source_idx}", src_desc)),
                        key=f"new_desc_{selected_part}_{source_idx}"
                    )

                st.divider()
                st.subheader("Specifications")
                col_tables, col_img = st.columns([1.5, 1.2])
                with col_tables:
                    st.markdown("**Specifications**")
                    meas_df_nova = pd.DataFrame(parse_measurements(source_iter.get('Specifications', source_iter.get('Constraints', ''))))
                    if meas_df_nova.empty or "Property" not in meas_df_nova.columns:
                        meas_df_nova = pd.DataFrame([{"Property": "", "Value": ""}])
                    edited_meas_nova = st.data_editor(meas_df_nova, num_rows="dynamic", use_container_width=True, key=f"new_meas_{selected_part}_{source_idx}")
                    
                with col_img:
                    pasted_img = render_image_input(
                        "Schematic",
                        key=f"new_img_{selected_part}",
                        current_image_path=source_iter.get('Image Path')
                    )

                # FEA Displacements and Screenshots for UP_01, HB_01 and brackets
                user_stiff_vals = {}
                user_lc_images = {}
                part_fea_cfg = get_part_fea_config(selected_part)
                part_lcs = part_fea_cfg.get("LC_KEYS", [])
                part_info = part_fea_cfg.get("LC_INFO", {})

                if part_lcs:
                    st.divider()
                    st.subheader("FEA Displacements")

                    ocr_vals_key = f"ocr_vals_new_{selected_part}_{source_idx}"
                    ocr_rev_key = f"ocr_rev_new_{selected_part}_{source_idx}"
                    if ocr_vals_key not in st.session_state:
                        st.session_state[ocr_vals_key] = {lc: None for lc in part_lcs}
                    if ocr_rev_key not in st.session_state:
                        st.session_state[ocr_rev_key] = 0

                    editor_widget_key = f"stiff_editor_new_{selected_part}_{source_idx}_{st.session_state[ocr_rev_key]}"
                    if editor_widget_key in st.session_state and isinstance(st.session_state[editor_widget_key], dict):
                        for r_idx_str, r_changes in st.session_state[editor_widget_key].get("edited_rows", {}).items():
                            r_idx = int(r_idx_str)
                            if 0 <= r_idx < len(part_lcs) and "Displacement (mm)" in r_changes:
                                st.session_state[ocr_vals_key][part_lcs[r_idx]] = r_changes["Displacement (mm)"]

                    part_decimals = part_fea_cfg.get("DECIMALS", 4)
                    if st.button("Paste FEA Print", key=f"btn_ocr_new_{selected_part}_{source_idx}", use_container_width=True):
                        pasted_fea = grab_image_from_clipboard()
                        if pasted_fea is not None:
                            det_lc, det_disp = extract_ntop_fea_from_image(pasted_fea, target_part=selected_part)
                            if det_lc and det_disp is not None and det_lc in part_lcs:
                                st.session_state[f"clipboard_new_{selected_part}_lc_{det_lc}"] = pasted_fea
                                st.session_state[ocr_vals_key][det_lc] = det_disp
                                st.session_state[ocr_rev_key] += 1
                                desc_lbl = part_info.get(det_lc, {}).get("desc", det_lc)
                                st.toast(f"{det_lc} ({desc_lbl}): {det_disp:.{part_decimals}f} mm")
                                st.rerun()
                            elif det_disp is not None:
                                st.warning(f"Detected {det_disp:.{part_decimals}f} mm, but Load Case title was not visible. Paste inside the specific LC tab below.")
                            else:
                                st.warning("Could not read nTop displacement scale from clipboard image.")
                        else:
                            st.warning("No image found on system clipboard. Capture a screenshot first (Win + Shift + S).")

                    editor_rows = []
                    for lc in part_lcs:
                        desc_str = part_info.get(lc, {}).get("desc", "")
                        lbl = f"{lc} ({desc_str})" if lc in ["LC8", "LC9"] and desc_str else lc
                        editor_rows.append({
                            "Load Case": lbl,
                            "Displacement (mm)": st.session_state[ocr_vals_key].get(lc, None),
                        })
                    df_stiff_input = pd.DataFrame(editor_rows)
                    df_stiff_input["Displacement (mm)"] = df_stiff_input["Displacement (mm)"].astype(float)

                    edited_stiff_df = st.data_editor(
                        df_stiff_input,
                        disabled=["Load Case"],
                        column_config={
                            "Displacement (mm)": st.column_config.NumberColumn(
                                "Displacement (mm)",
                                min_value=0.0,
                                max_value=10.0,
                                step=10 ** (-part_decimals),
                                format=f"%.{part_decimals}f"
                            )
                        },
                        use_container_width=True,
                        key=editor_widget_key
                    )

                    for r in edited_stiff_df.to_dict('records'):
                        lc_id = r["Load Case"].split(" ")[0]
                        raw_d = r.get("Displacement (mm)", None)
                        try:
                            if raw_d is None or (isinstance(raw_d, float) and math.isnan(raw_d)):
                                st.session_state[ocr_vals_key][lc_id] = None
                                user_stiff_vals[lc_id] = 0.0
                            else:
                                num_d = float(raw_d)
                                st.session_state[ocr_vals_key][lc_id] = None if math.isnan(num_d) else num_d
                                user_stiff_vals[lc_id] = 0.0 if math.isnan(num_d) else num_d
                        except Exception:
                            user_stiff_vals[lc_id] = 0.0

                    def make_new_lc_ocr_cb(target_lc):
                        def _cb(img):
                            _, det_disp = extract_ntop_fea_from_image(img, target_part=selected_part)
                            if det_disp is not None:
                                st.session_state[ocr_vals_key][target_lc] = det_disp
                                st.session_state[ocr_rev_key] += 1
                                st.toast(f"{target_lc}: {det_disp:.{part_decimals}f} mm")
                            else:
                                st.toast("Image pasted from clipboard.")
                        return _cb

                    with st.expander("FEA Screenshots", expanded=False):
                        lc_tabs = st.tabs([f"{lc}" for lc in part_lcs])
                        for idx_lc, lc in enumerate(part_lcs):
                            with lc_tabs[idx_lc]:
                                desc_lbl = part_info.get(lc, {}).get("desc", lc)
                                st.write(f"**{lc} - {desc_lbl}**")
                                pasted_lc_img = render_image_input(
                                    label=f"{lc} Screenshot",
                                    key=f"new_{selected_part}_lc_{lc}",
                                    ocr_callback=make_new_lc_ocr_cb(lc)
                                )
                                if pasted_lc_img is not None:
                                    user_lc_images[lc] = pasted_lc_img

                elif is_structural:
                    st.divider()
                    st.subheader("Structural FEA")
                    src_struct = parse_structural_data(source_iter.get("Stiffness Data", ""))
                    with st.expander("Structural FEA", expanded=True):
                        c_f1, c_f2 = st.columns(2)
                        with c_f1:
                            def_lc = src_struct.get("load_case", "Wishbone Lateral 4.4 kN" if selected_part in ["UP_02", "UpperSupport Version"] else "Bump / Cornering")
                            lc_input = st.text_input("Critical Load Case", value=def_lc, key=f"new_struct_lc_{selected_part}_{source_idx}")
                            mat_options = ["Al 7075-T6 (Sy = 503 MPa)", "Al 6082-T6 (Sy = 260 MPa)", "Ti-6Al-4V (Sy = 880 MPa)", "Custom"]
                            mat_choice = st.selectbox("Material", mat_options, index=0, key=f"new_struct_mat_{selected_part}_{source_idx}")
                            yield_val = 503.0 if "7075" in mat_choice else (260.0 if "6082" in mat_choice else (880.0 if "Ti" in mat_choice else 500.0))
                        with c_f2:
                            def_stress = float(src_struct.get("stress", 0.0))
                            stress_input = st.number_input("Peak Stress (MPa)", min_value=0.0, value=def_stress, format="%.1f", key=f"new_stress_{selected_part}_{source_idx}")
                            auto_sf = round(yield_val / stress_input, 2) if stress_input > 0 else float(src_struct.get("sf", 0.0))
                            sf_input = st.number_input("Safety Factor", min_value=0.0, value=auto_sf, format="%.2f", key=f"new_sf_{selected_part}_{source_idx}")
                            def_disp = float(src_struct.get("disp", 0.0))
                            disp_input = st.number_input("Max Deflection (mm)", min_value=0.0, value=def_disp, format="%.4f", key=f"new_disp_{selected_part}_{source_idx}")

                        pasted_fea_img = render_image_input(
                            label="FEA Screenshot",
                            key=f"new_bracket_fea_{selected_part}"
                        )
                        if pasted_fea_img is not None:
                            user_lc_images["FEA"] = pasted_fea_img
                        if stress_input > 0 or disp_input > 0:
                            user_stiff_vals = {
                                "load_case": lc_input.strip(),
                                "stress": stress_input,
                                "yield_strength": yield_val,
                                "sf": sf_input,
                                "disp": disp_input,
                                "material": mat_choice.split(" (")[0]
                            }

                st.divider()
                if st.button("Save Iteration", type="primary", use_container_width=True, key=f"btn_save_new_{selected_part}"):
                    new_version = str(new_version).lstrip('vV').strip()
                    img_path = save_pasted_image(pasted_img, selected_part, new_version)
                    if img_path:
                        st.session_state.pop(f"clipboard_new_img_{selected_part}", None)
                    else:
                        src_img = str(source_iter.get("Image Path", "")).strip()
                        img_path = "" if src_img == "nan" else src_img

                    saved_lc_paths = {}
                    for lc, img in user_lc_images.items():
                        s_path = save_lc_image(img, selected_part, new_version, lc)
                        if s_path:
                            saved_lc_paths[lc] = s_path
                            st.session_state.pop(f"clipboard_new_up_lc_{lc}", None)
                            st.session_state.pop(f"clipboard_new_bracket_fea_{selected_part}", None)

                    final_meas = [r for r in edited_meas_nova.to_dict('records') if str(r.get('Property', '')).strip() != ""]
                    comp_tol = str(latest_iter.get("Tolerancing Data", "")).strip()
                    if comp_tol in ["nan", "[]"]:
                        comp_tol = ""

                    comp_just = str(latest_iter.get("Justification", "")).strip()
                    if comp_just == "nan":
                        comp_just = ""

                    text_cols_save = ["Description", "Specifications", "Tolerancing Data", "Stiffness Data", "LC Images", "Justification", "Image Path"]
                    for col_name in text_cols_save:
                        if col_name in df_parts.columns and df_parts[col_name].dtype != object:
                            df_parts[col_name] = df_parts[col_name].astype(object)

                    new_row = {
                        "Subassembly": latest_iter['Subassembly'],
                        "Part Name": selected_part,
                        "Version": new_version,
                        "Date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "Weight (g)": weight,
                        "Description": desc,
                        "Specifications": format_measurements(final_meas),
                        "Tolerancing Data": comp_tol,
                        "Image Path": img_path,
                        "Stiffness Data": format_stiffness_data(user_stiff_vals, lc_keys=part_lcs, decimals=part_decimals) if part_lcs else (json.dumps(user_stiff_vals) if user_stiff_vals else ""),
                        "LC Images": format_lc_images(saved_lc_paths),
                        "Justification": comp_just
                    }
                    # Prevent duplicate (Part Name, Version) rows
                    dup_mask = (df_parts["Part Name"].astype(str) == str(selected_part)) & (df_parts["Version"].astype(str) == str(new_version))
                    if dup_mask.any():
                        df_parts = df_parts[~dup_mask].reset_index(drop=True)
                    df_parts = pd.concat([df_parts, pd.DataFrame([new_row])], ignore_index=True)
                    part_mask = df_parts["Part Name"] == selected_part
                    df_parts.loc[part_mask, "Justification"] = comp_just
                    df_parts.loc[part_mask, "Tolerancing Data"] = comp_tol
                    save_parts(df_parts)
                    st.session_state.pop(f"ocr_vals_new_{selected_part}_{source_idx}", None)
                    st.session_state.pop(f"ocr_rev_new_{selected_part}_{source_idx}", None)
                    new_idx = df_parts.index[-1]
                    st.session_state[f"pending_edit_{selected_part}"] = new_idx
                    st.session_state[f"pending_eval_{selected_part}"] = str(new_version)
                    st.success(f"Iteration {format_v_label(new_version)} created.")
                    st.rerun()

            # ---------- TAB: EDIT ----------
            with tab_editar:
                st.subheader("Edit Iteration" if not is_oem else "Edit Part")
                version_indices = part_data.sort_values("Date", ascending=False).index.tolist()
                def format_version_option(idx):
                    r = part_data.loc[idx]
                    return f"{format_v_label(r['Version'])} ({r['Date']})"

                sel_key = f"sel_edit_{selected_part}"
                if sel_key in st.session_state and st.session_state[sel_key] not in version_indices:
                    st.session_state.pop(sel_key, None)

                selected_idx = st.selectbox(
                    "Select Version",
                    version_indices,
                    format_func=format_version_option,
                    index=0,
                    key=sel_key
                )
                row_to_edit = part_data.loc[selected_idx]
                version_to_edit = row_to_edit["Version"]

                col1, col2 = st.columns(2)
                with col1:
                    edit_weight = st.number_input(
                        "Weight (g)",
                        value=float(row_to_edit['Weight (g)']),
                        format="%.2f",
                        key=f"edit_wt_{selected_idx}"
                    )
                with col2:
                    edit_desc = st.text_area(
                        "Description",
                        value=str(row_to_edit.get('Description', '')),
                        height=calc_description_height(row_to_edit.get('Description', '')),
                        key=f"edit_desc_{selected_idx}"
                    )

                st.divider()
                st.subheader("Specifications")
                col_tables, col_img = st.columns([1.5, 1.2])
                with col_tables:
                    st.markdown("**Specifications**")
                    meas_df_edit = pd.DataFrame(parse_measurements(row_to_edit.get('Specifications', row_to_edit.get('Constraints', ''))))
                    if meas_df_edit.empty or "Property" not in meas_df_edit.columns:
                        meas_df_edit = pd.DataFrame([{"Property": "", "Value": ""}])
                    edited_meas_edit = st.data_editor(
                        meas_df_edit,
                        num_rows="dynamic",
                        use_container_width=True,
                        key=f"de_edit_meas_{selected_idx}"
                    )

                with col_img:
                    pasted_edit_img = render_image_input(
                        "Schematic",
                        key=f"edit_img_{selected_idx}",
                        current_image_path=row_to_edit.get('Image Path')
                    )

                # FEA editing for UP_01, HB_01 and brackets
                user_edit_stiff_vals = {}
                updated_lc_images_pasted = {}
                curr_lc_images = parse_lc_images(row_to_edit.get("LC Images", ""))
                curr_struct = parse_structural_data(row_to_edit.get("Stiffness Data", ""))
                part_fea_cfg = get_part_fea_config(selected_part)
                part_lcs = part_fea_cfg.get("LC_KEYS", [])
                part_info = part_fea_cfg.get("LC_INFO", {})

                if part_lcs:
                    st.divider()
                    st.subheader("FEA Displacements")

                    curr_edit_stiff = parse_stiffness_data(row_to_edit.get("Stiffness Data", ""), lc_keys=part_lcs)
                    if not curr_edit_stiff:
                        if "FST15" in str(version_to_edit).upper():
                            curr_edit_stiff = part_fea_cfg.get("DEFAULT_FST15", {})
                        else:
                            curr_edit_stiff = part_fea_cfg.get("DEFAULT_FST16", {})

                    ocr_edit_vals_key = f"ocr_vals_edit_{selected_idx}"
                    ocr_edit_rev_key = f"ocr_rev_edit_{selected_idx}"
                    if ocr_edit_vals_key not in st.session_state:
                        st.session_state[ocr_edit_vals_key] = {
                            lc: curr_edit_stiff.get(lc, part_info.get(lc, {}).get("default_fst16", 0.0)) for lc in part_lcs
                        }
                    if ocr_edit_rev_key not in st.session_state:
                        st.session_state[ocr_edit_rev_key] = 0

                    edit_widget_key = f"stiff_editor_edit_{selected_idx}_{st.session_state[ocr_edit_rev_key]}"
                    if edit_widget_key in st.session_state and isinstance(st.session_state[edit_widget_key], dict):
                        for r_idx_str, r_changes in st.session_state[edit_widget_key].get("edited_rows", {}).items():
                            r_idx = int(r_idx_str)
                            if 0 <= r_idx < len(part_lcs) and "Displacement (mm)" in r_changes:
                                st.session_state[ocr_edit_vals_key][part_lcs[r_idx]] = r_changes["Displacement (mm)"]

                    part_decimals = part_fea_cfg.get("DECIMALS", 4)
                    if st.button("Paste FEA Print", key=f"btn_ocr_edit_{selected_idx}", use_container_width=True):
                        pasted_fea_edit = grab_image_from_clipboard()
                        if pasted_fea_edit is not None:
                            det_lc, det_disp = extract_ntop_fea_from_image(pasted_fea_edit, target_part=selected_part)
                            if det_lc and det_disp is not None and det_lc in part_lcs:
                                st.session_state[f"clipboard_edit_{selected_part}_lc_{selected_idx}_{det_lc}"] = pasted_fea_edit
                                st.session_state[ocr_edit_vals_key][det_lc] = det_disp
                                st.session_state[ocr_edit_rev_key] += 1
                                desc_lbl = part_info.get(det_lc, {}).get("desc", det_lc)
                                st.toast(f"{det_lc} ({desc_lbl}): {det_disp:.{part_decimals}f} mm")
                                st.rerun()
                            elif det_disp is not None:
                                st.warning(f"Detected {det_disp:.{part_decimals}f} mm, but Load Case title was not visible. Paste inside the specific LC tab below.")
                            else:
                                st.warning("Could not read nTop displacement scale from clipboard image.")
                        else:
                            st.warning("No image found on system clipboard. Capture a screenshot first (Win + Shift + S).")

                    editor_rows_edit = []
                    for lc in part_lcs:
                        def_val = part_info.get(lc, {}).get("default_fst16", 0.0)
                        val = st.session_state[ocr_edit_vals_key].get(lc, curr_edit_stiff.get(lc, def_val))
                        desc_str = part_info.get(lc, {}).get("desc", "")
                        lbl = f"{lc} ({desc_str})" if lc in ["LC8", "LC9"] and desc_str else lc
                        editor_rows_edit.append({
                            "Load Case": lbl,
                            "Displacement (mm)": val,
                        })
                    df_stiff_edit = pd.DataFrame(editor_rows_edit)
                    df_stiff_edit["Displacement (mm)"] = df_stiff_edit["Displacement (mm)"].astype(float)

                    edited_stiff_df_edit = st.data_editor(
                        df_stiff_edit,
                        disabled=["Load Case"],
                        column_config={
                            "Displacement (mm)": st.column_config.NumberColumn(
                                "Displacement (mm)",
                                min_value=0.0,
                                max_value=10.0,
                                step=10 ** (-part_decimals),
                                format=f"%.{part_decimals}f"
                            )
                        },
                        use_container_width=True,
                        key=edit_widget_key
                    )

                    for r in edited_stiff_df_edit.to_dict('records'):
                        lc_id = r["Load Case"].split(" ")[0]
                        raw_d = r.get("Displacement (mm)", None)
                        try:
                            if raw_d is None or (isinstance(raw_d, float) and math.isnan(raw_d)):
                                st.session_state[ocr_edit_vals_key][lc_id] = None
                                user_edit_stiff_vals[lc_id] = 0.0
                            else:
                                num_d = float(raw_d)
                                clean_d = 0.0 if math.isnan(num_d) else num_d
                                st.session_state[ocr_edit_vals_key][lc_id] = None if math.isnan(num_d) else num_d
                                user_edit_stiff_vals[lc_id] = clean_d
                        except Exception:
                            user_edit_stiff_vals[lc_id] = 0.0

                    def make_edit_lc_ocr_cb(target_lc):
                        def _cb(img):
                            _, det_disp = extract_ntop_fea_from_image(img, target_part=selected_part)
                            if det_disp is not None:
                                st.session_state[ocr_edit_vals_key][target_lc] = det_disp
                                st.session_state[ocr_edit_rev_key] += 1
                                st.toast(f"{target_lc}: {det_disp:.{part_decimals}f} mm")
                            else:
                                st.toast("Image pasted from clipboard.")
                        return _cb

                    with st.expander("FEA Screenshots", expanded=False):
                        lc_tabs_edit = st.tabs([f"{lc}" for lc in part_lcs])
                        for idx_lc, lc in enumerate(part_lcs):
                            with lc_tabs_edit[idx_lc]:
                                desc_lbl = part_info.get(lc, {}).get("desc", lc)
                                st.write(f"**{lc} - {desc_lbl}**")
                                cur_lc_p = curr_lc_images.get(lc, None)
                                p_edit_lc = render_image_input(
                                    label=f"{lc} Screenshot",
                                    key=f"edit_{selected_part}_lc_{selected_idx}_{lc}",
                                    current_image_path=cur_lc_p,
                                    ocr_callback=make_edit_lc_ocr_cb(lc)
                                )
                                if p_edit_lc is not None:
                                    updated_lc_images_pasted[lc] = p_edit_lc

                elif is_structural:
                    st.divider()
                    st.subheader("Structural FEA")
                    curr_struct = parse_structural_data(row_to_edit.get("Stiffness Data", ""))
                    exp_edit = bool(curr_struct)
                    with st.expander("Structural FEA", expanded=exp_edit):
                        c_f1, c_f2 = st.columns(2)
                        with c_f1:
                            cur_lc = curr_struct.get("load_case", "Wishbone Lateral 4.4 kN" if selected_part == "UpperSupport Version" else "Bump / Cornering")
                            edit_lc_input = st.text_input("Critical Load Case", value=cur_lc, key=f"edit_struct_lc_{selected_idx}")
                            cur_mat_str = curr_struct.get("material", "Al 7075-T6")
                            mat_options = ["Al 7075-T6 (Sy = 503 MPa)", "Al 6082-T6 (Sy = 260 MPa)", "Ti-6Al-4V (Sy = 880 MPa)", "Custom"]
                            mat_idx = 0
                            if "6082" in cur_mat_str: mat_idx = 1
                            elif "Ti" in cur_mat_str: mat_idx = 2
                            edit_mat = st.selectbox("Material", mat_options, index=mat_idx, key=f"edit_struct_mat_{selected_idx}")
                            edit_yield_val = 503.0 if "7075" in edit_mat else (260.0 if "6082" in edit_mat else (880.0 if "Ti" in edit_mat else 500.0))
                        with c_f2:
                            edit_stress = st.number_input("Peak Stress (MPa)", min_value=0.0, value=float(curr_struct.get("stress", 0.0)), format="%.1f", key=f"edit_stress_{selected_idx}")
                            edit_auto_sf = round(edit_yield_val / edit_stress, 2) if edit_stress > 0 else float(curr_struct.get("sf", 0.0))
                            edit_sf = st.number_input("Safety Factor", min_value=0.0, value=float(edit_auto_sf), format="%.2f", key=f"edit_sf_{selected_idx}")
                            edit_disp = st.number_input("Max Deflection (mm)", min_value=0.0, value=float(curr_struct.get("disp", 0.0)), format="%.4f", key=f"edit_disp_{selected_idx}")

                        existing_fea_img = curr_lc_images.get("FEA", curr_lc_images.get("Structural", None))
                        pasted_edit_fea = render_image_input(
                            label="FEA Screenshot",
                            key=f"edit_bracket_fea_{selected_idx}",
                            current_image_path=existing_fea_img
                        )
                        if pasted_edit_fea is not None:
                            updated_lc_images_pasted["FEA"] = pasted_edit_fea
                        if edit_stress > 0 or edit_disp > 0:
                            user_edit_stiff_vals = {
                                "load_case": edit_lc_input.strip(),
                                "stress": edit_stress,
                                "yield_strength": edit_yield_val,
                                "sf": edit_sf,
                                "disp": edit_disp,
                                "material": edit_mat.split(" (")[0]
                            }

                st.divider()
                if len(part_data) > 1 and not is_oem:
                    c_save, c_del_iter, c_del = st.columns([2, 2, 2])
                    with c_save:
                        save_clicked = st.button("Save Changes", type="primary", use_container_width=True, key=f"save_edit_btn_{selected_idx}")
                    with c_del_iter:
                        del_iter_clicked = st.button("Delete Iteration", use_container_width=True, key=f"del_iter_btn_{selected_idx}")
                    with c_del:
                        del_clicked = st.button("Delete Part", use_container_width=True, key=f"del_part_btn_{selected_idx}")
                else:
                    c_save, c_del = st.columns([3, 2])
                    with c_save:
                        save_clicked = st.button("Save Changes", type="primary", use_container_width=True, key=f"save_edit_btn_{selected_idx}")
                    with c_del:
                        del_clicked = st.button("Delete Part", use_container_width=True, key=f"del_part_btn_{selected_idx}")
                    del_iter_clicked = False

                if save_clicked:
                    if pasted_edit_img is not None:
                        new_img_path = save_pasted_image(pasted_edit_img, selected_part, version_to_edit)
                        if new_img_path:
                            df_parts.at[selected_idx, 'Image Path'] = new_img_path
                            st.session_state.pop(f"clipboard_edit_img_{selected_idx}", None)

                    final_lc_images = dict(curr_lc_images)
                    for lc, img in updated_lc_images_pasted.items():
                        new_lc_p = save_lc_image(img, selected_part, version_to_edit, lc)
                        if new_lc_p:
                            final_lc_images[lc] = new_lc_p
                            st.session_state.pop(f"clipboard_edit_up_lc_{selected_idx}_{lc}", None)
                            st.session_state.pop(f"clipboard_edit_bracket_fea_{selected_idx}", None)

                    final_meas = [r for r in edited_meas_edit.to_dict('records') if str(r.get('Property', '')).strip() != ""]

                    # Ensure text columns have object dtype so string assignment never raises LossySetitemError
                    text_cols_save = ["Description", "Specifications", "Tolerancing Data", "Stiffness Data", "LC Images", "Justification", "Image Path"]
                    for col_name in text_cols_save:
                        if col_name in df_parts.columns and df_parts[col_name].dtype != object:
                            df_parts[col_name] = df_parts[col_name].astype(object)

                    df_parts.at[selected_idx, 'Weight (g)'] = edit_weight
                    df_parts.at[selected_idx, 'Description'] = edit_desc
                    df_parts.at[selected_idx, 'Specifications'] = format_measurements(final_meas)
                    if part_lcs:
                        df_parts.at[selected_idx, 'Stiffness Data'] = format_stiffness_data(user_edit_stiff_vals, lc_keys=part_lcs, decimals=part_decimals)
                        df_parts.at[selected_idx, 'LC Images'] = format_lc_images(final_lc_images)
                    else:
                        if user_edit_stiff_vals:
                            df_parts.at[selected_idx, 'Stiffness Data'] = json.dumps(user_edit_stiff_vals)
                        elif curr_struct:
                            df_parts.at[selected_idx, 'Stiffness Data'] = json.dumps(curr_struct)
                        if final_lc_images:
                            df_parts.at[selected_idx, 'LC Images'] = format_lc_images(final_lc_images)

                    save_parts(df_parts)
                    st.success(f"Iteration {format_v_label(version_to_edit)} updated.")
                    st.rerun()

                if del_iter_clicked:
                    if len(part_data) <= 1:
                        st.error("Cannot delete only iteration.")
                    else:
                        df_parts = df_parts.drop(selected_idx)
                        save_parts(df_parts)
                        st.session_state.pop(sel_key, None)
                        st.success("Iteration deleted.")
                        st.rerun()

                if del_clicked:
                    part_to_trash = df_parts[df_parts["Part Name"] == selected_part]
                    df_trash = pd.concat([df_trash, part_to_trash], ignore_index=True)
                    save_trash(df_trash)
                    df_parts = df_parts[df_parts["Part Name"] != selected_part]
                    save_parts(df_parts)
                    st.session_state.pop(sel_key, None)
                    st.session_state.current_page = "Wheel Assembly"
                    st.session_state.selected_part = None
                    st.rerun()
