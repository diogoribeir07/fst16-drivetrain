"""
Comprehensive unit test suite for Bearing Sizer & Selector and FST16_Wheel_Bearing_Calculator.xlsx.
"""

import os
import re
import io
import math
import zipfile
import openpyxl
import pandas as pd

from bearing_sizer import (
    BEARING_DATA,
    get_bearing_catalog_df,
    get_bearing_dict,
    compute_kinematics_and_loads,
    evaluate_all_load_cases,
    read_bytes_safely,
    EXCEL_CALC_PATH
)

def get_test_workbook():
    raw = read_bytes_safely(EXCEL_CALC_PATH)
    assert raw is not None, f"Could not read {EXCEL_CALC_PATH}"
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    return wb, raw

def test_excel_file_exists_and_sheets():
    assert os.path.exists(EXCEL_CALC_PATH), f"Excel file {EXCEL_CALC_PATH} does not exist"
    wb, _ = get_test_workbook()
    assert "Bearing_Database" in wb.sheetnames
    assert "Bearing_Calculator" in wb.sheetnames

def test_excel_database_contents():
    _, raw = get_test_workbook()
    df_db = pd.read_excel(io.BytesIO(raw), sheet_name="Bearing_Database", header=2)
    assert len(df_db) >= 10, f"Expected 10 candidate bearings, found {len(df_db)}"
    required_cols = [
        "Part_Number", "Type", "Bore_d_mm", "OD_D_mm", "Width_B_mm", "Pitch_dm_mm",
        "Dynamic_Cr_kN", "Static_C0r_kN", "Contact_Angle_deg", "Ball_Diam_Dw_mm",
        "Ball_Count_Z", "Limiting_Speed_Grease_rpm", "Limiting_Speed_Oil_rpm",
        "Radial_Stiffness_Kr_kN_mm", "Axial_Stiffness_Ka_kN_mm", "Mass_g",
        "Unit_Cost_EUR", "Lead_Time_Weeks", "Status"
    ]
    for col in required_cols:
        assert col in df_db.columns, f"Column {col} missing in Bearing_Database sheet"
    
    # Verify primary candidate bearings
    part_numbers = df_db["Part_Number"].tolist()
    assert "Schaeffler 61816-Y" in part_numbers
    assert "FAG B71816-E-TPA-P4S-UL" in part_numbers
    assert "FAG B71816-E-TPA-P4S-UM" in part_numbers
    assert "FAG B71816-E-TPA-P4S-UH" in part_numbers
    assert "FAG HCB71816-E-TPA-P4S-UL" in part_numbers
    assert "Schaeffler 61815-Y" in part_numbers
    assert "FAG B71815-E-TPA-P4S-UL" in part_numbers

def test_excel_datavalidation_and_formula_integrity():
    wb, raw = get_test_workbook()
    ws = wb["Bearing_Calculator"]
    
    # Verify DataValidation exists (in openpyxl dataValidation or OpenXML extLst)
    has_dv = False
    if len(ws.data_validations.dataValidation) > 0:
        dv = ws.data_validations.dataValidation[0]
        assert "Bearing_Database" in dv.formula1, f"DataValidation formula does not reference Bearing_Database: {dv.formula1}"
        has_dv = True
    else:
        zf = zipfile.ZipFile(io.BytesIO(raw))
        for name in zf.namelist():
            if "worksheets" in name:
                xml = zf.read(name).decode("utf-8", errors="ignore")
                if "datavalidation" in xml.lower() and "bearing_database" in xml.lower():
                    has_dv = True
                    break
    assert has_dv, "No DataValidation referencing Bearing_Database found in Bearing_Calculator!"
    
    # Check Inner S0 formula (row with S0 = C0r / P0)
    s0_inner_found = False
    l10h_inner_found = False
    frout_found = False
    for r in range(1, ws.max_row + 1):
        lbl = ws.cell(row=r, column=2).value
        if lbl and "Static Safety Factor" in str(lbl):
            val_in = str(ws.cell(row=r, column=4).value)
            # Must divide by D{p0}, not C{p0}
            assert "/D" in val_in, f"Inner S0 formula in row {r} does not divide by column D: {val_in}"
            s0_inner_found = True
        elif lbl and "Basic Rating Fatigue Life" in str(lbl):
            val_in = str(ws.cell(row=r, column=4).value)
            assert "/D" in val_in, f"Inner L10h formula in row {r} does not divide by column D: {val_in}"
            l10h_inner_found = True
        elif lbl and "Peak Outer Radial Force" in str(lbl):
            val_fro = str(ws.cell(row=r, column=3).value)
            assert "/2" in val_fro, f"Peak Outer Radial Force formula does not divide Fz by 2: {val_fro}"
            frout_found = True
            
    assert s0_inner_found, "Static Safety Factor row not found in Bearing_Calculator"
    assert l10h_inner_found, "Basic Rating Fatigue Life row not found in Bearing_Calculator"
    assert frout_found, "Peak Outer Radial Force row not found in Bearing_Calculator"

def test_excel_no_forbidden_terms_or_emojis():
    wb, _ = get_test_workbook()
    forbidden = [
        "Component Status", "DL", "sign-off", "signoff", "Approved Components",
        "Under Trade-Off Study", "Ready for DL Review"
    ]
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    sc = str(cell)
                    for fb in forbidden:
                        if fb == "DL":
                            assert not re.search(r'\bDL\b', sc), f"Forbidden term DL found in {sname}: {sc}"
                        else:
                            assert fb.lower() not in sc.lower(), f"Forbidden term {fb} found in {sname}: {sc}"
                    for c in sc:
                        assert not (ord(c) > 127 and (ord(c) >= 126976 or (9728 <= ord(c) <= 10175))), f"Emoji found in {sname}: {sc}"

def test_dual_61816_physics_calculation():
    res = compute_kinematics_and_loads(
        outer_pname="Schaeffler 61816-Y",
        inner_pname="Schaeffler 61816-Y",
        fy=4436.26,
        fz=1787.31,
        reff=0.250,
        lb=32.0,
        v_kmh=100.0,
        temp_c=80.0
    )
    # Check overturning moment
    assert abs(res["mx"] - 1109.065) < 0.1
    # Check effective spread for deep groove
    assert abs(res["effective_spread_h"] - 32.0) < 0.01
    # Check couple
    assert abs(res["f_couple"] - 34658.28) < 1.0
    # Peak outer radial load Fr_out = F_couple + Fz/2
    assert abs(res["fr_out"] - 35551.94) < 1.0
    # Static safety factor S0
    assert 0.35 <= res["s0_out"] <= 0.42
    # Motosh torque (around 11 N*m)
    assert 9.0 <= res["torque_nm"] <= 13.0
    # Thermal bore fit change (expanded loose fit)
    assert res["net_fit_at_temp_mm"] < 0.0

    # Physical units check on Camber Stiffness and Deflection
    assert 4000.0 <= res["c_tilt_deg"] <= 5000.0, f"Expected c_tilt_deg around 4468 N*m/deg, got {res['c_tilt_deg']}"
    assert 0.20 <= res["gamma_deg"] <= 0.30, f"Expected realistic camber deflection around 0.25 deg, got {res['gamma_deg']}"
    assert 0.8 <= res["delta_tire_top_mm"] <= 1.5, f"Expected tire top deflection around 1.1 mm, got {res['delta_tire_top_mm']}"

def test_dual_71816_angular_contact_upgrade():
    res = compute_kinematics_and_loads(
        outer_pname="FAG B71816-E-TPA-P4S-UL",
        inner_pname="FAG B71816-E-TPA-P4S-UL",
        fy=4436.26,
        fz=1787.31,
        reff=0.250,
        lb=32.0,
        v_kmh=100.0,
        temp_c=80.0
    )
    # Apex distance a_out
    assert abs(res["a_out"] - 20.98) < 0.1
    # Effective spread H ~ 73.97 mm
    assert abs(res["effective_spread_h"] - 73.97) < 0.2
    # Overturning couple reduced by ~57%
    assert abs(res["f_couple"] - 14993.8) < 50.0
    # Static safety factor S0 single ~ 0.98
    assert 0.90 <= res["s0_out"] <= 1.05
    # Camber stiffness gain ratio vs 32mm physical span
    assert res["stiffness_ratio_vs_physical"] > 5.0
    # Motosh torque for 84 N light preload
    assert 1.0 <= res["torque_nm"] <= 2.0
    
    # Camber rigidity upgrade (> 20,000 N*m/deg, gamma ~ 0.05 deg)
    assert res["c_tilt_deg"] > 20000.0, f"Expected c_tilt_deg > 20,000 N*m/deg, got {res['c_tilt_deg']}"
    assert 0.04 <= res["gamma_deg"] <= 0.07, f"Expected gamma around 0.05 deg, got {res['gamma_deg']}"
    assert 0.15 <= res["delta_tire_top_mm"] <= 0.35, f"Expected tire top deflection ~ 0.23 mm, got {res['delta_tire_top_mm']}"

def test_kinematics_gyroscopic_moment_and_antiskid():
    res_ac = compute_kinematics_and_loads("FAG B71816-E-TPA-P4S-UL", "FAG B71816-E-TPA-P4S-UL", 4436.26, 1787.31, 0.250, 32.0, 100.0, 80.0)
    # Gyroscopic moment Mg > 0 for 25 deg ACBB
    assert res_ac["mg_ball_nmm"] > 0.0, f"Expected positive Mg, got {res_ac['mg_ball_nmm']}"
    # Anti-skid preload threshold > 0
    assert res_ac["fa_skid_n"] > 0.0, f"Expected positive fa_skid_n, got {res_ac['fa_skid_n']}"
    # Centrifugal force Fc > 0
    assert res_ac["fc_ball"] > 0.0

    res_dg = compute_kinematics_and_loads("Schaeffler 61816-Y", "Schaeffler 61816-Y", 4436.26, 1787.31, 0.250, 32.0, 100.0, 80.0)
    # For deep groove, alpha = 0 -> Mg = 0
    assert res_dg["mg_ball_nmm"] == 0.0
    # Minimum radial load threshold = 0.01 * C0r = 138 N
    assert abs(res_dg["fa_skid_n"] - 138.0) < 1.0

def test_preload_classes_um_uh():
    res_um = compute_kinematics_and_loads("FAG B71816-E-TPA-P4S-UM", "FAG B71816-E-TPA-P4S-UM", 4436.26, 1787.31, 0.250, 32.0, 100.0, 80.0)
    assert res_um["f_preload"] == 302.0
    assert 3.5 <= res_um["torque_nm"] <= 4.5

    res_uh = compute_kinematics_and_loads("FAG B71816-E-TPA-P4S-UH", "FAG B71816-E-TPA-P4S-UH", 4436.26, 1787.31, 0.250, 32.0, 100.0, 80.0)
    assert res_uh["f_preload"] == 712.0
    assert 8.5 <= res_uh["torque_nm"] <= 10.0

def test_zero_speed_division_by_zero_guard():
    # Should not raise ZeroDivisionError
    res_zero = compute_kinematics_and_loads("Schaeffler 61816-Y", "Schaeffler 61816-Y", 4436.26, 1787.31, 0.250, 32.0, 0.0, 80.0)
    assert res_zero["n_wheel"] == 0.0
    assert res_zero["l10h_out"] >= 99999.0
    assert res_zero["kappa"] >= 90.0

def test_all_8_load_cases_evaluation():
    df_lc = evaluate_all_load_cases(
        outer_pname="FAG B71816-E-TPA-P4S-UL",
        effective_spread_h=73.97,
        reff=0.250
    )
    assert len(df_lc) == 8
    cases = df_lc["Load Case"].tolist()
    assert any("LC1" in c for c in cases)
    assert any("LC2" in c for c in cases)
    assert any("LC3" in c for c in cases)
    assert any("LC4" in c for c in cases)
    assert any("LC5" in c for c in cases)
    assert any("LC6" in c for c in cases)
    assert any("LC7" in c for c in cases)
    assert any("Bump" in c for c in cases)

    # In LC1, LC2, LC Bump, S0 should be high (>= 4.0)
    lc1_s0 = df_lc[df_lc["Load Case"].str.contains("LC1")]["Static S0"].iloc[0]
    assert lc1_s0 >= 4.0
    lc_bump_s0 = df_lc[df_lc["Load Case"].str.contains("Bump")]["Static S0"].iloc[0]
    assert lc_bump_s0 >= 4.0

def test_ui_brevity_and_cleanliness():
    with open(os.path.join(os.path.dirname(__file__), "bearing_sizer.py"), "r", encoding="utf-8") as f:
        src = f.read()
    
    # Zero emojis
    emojis = [c for c in src if ord(c) > 127 and (ord(c) >= 126976 or (9728 <= ord(c) <= 10175))]
    assert len(emojis) == 0, f"Found emojis: {emojis}"
    
    # Zero forbidden terms
    forbidden = [
        "Component Status", "DEFAULT_STATUSES", "DL Review", "DL Status",
        "DL Decision", "DL Recommendation", "Component Trade-Off Study",
        "sign-off", "Approved Components", "Under Trade-Off Study",
        "Ready for DL Review", "Ready for Review", "EXCEL_TASKS", "load_design_tasks"
    ]
    for term in forbidden:
        assert term.lower() not in src.lower(), f"Found forbidden term: {term}"
    assert not re.search(r'\bDL\b', src), "Found standalone DL in bearing_sizer.py"

    # UI element word count check (<= 3 words)
    patterns = [
        (r'st\.title\(f?[\'"]([^\'"]+)[\'"]\)', 'st.title'),
        (r'st\.subheader\(f?[\'"]([^\'"]+)[\'"]\)', 'st.subheader'),
        (r'st\.tabs\(\[([^\]]+)\]\)', 'st.tabs'),
        (r'st\.button\(f?[\'"]([^\'"]+)[\'"]', 'st.button'),
        (r'st\.download_button\(f?label=[\'"]([^\'"]+)[\'"]', 'st.download_button'),
        (r'st\.expander\(f?[\'"]([^\'"]+)[\'"]', 'st.expander'),
        (r'(?:\w+\.)?metric\(f?[\'"]([^\'"]+)[\'"]', 'metric'),
        (r'st\.selectbox\(f?[\'"]([^\'"]+)[\'"]', 'st.selectbox'),
        (r'st\.number_input\(f?[\'"]([^\'"]+)[\'"]', 'st.number_input'),
        (r'st\.slider\(f?[\'"]([^\'"]+)[\'"]', 'st.slider'),
    ]
    for pat, kind in patterns:
        for m in re.finditer(pat, src):
            text = m.group(1).strip()
            if kind == 'st.tabs':
                for item in text.split(','):
                    words = item.strip(' \n\'"').split()
                    assert len(words) <= 3, f"Tab '{item}' has {len(words)} words (>3 words)"
            else:
                words = text.split()
                assert len(words) <= 3, f"{kind} '{text}' has {len(words)} words (>3 words)"

if __name__ == "__main__":
    test_excel_file_exists_and_sheets()
    test_excel_database_contents()
    test_excel_datavalidation_and_formula_integrity()
    test_excel_no_forbidden_terms_or_emojis()
    test_dual_61816_physics_calculation()
    test_dual_71816_angular_contact_upgrade()
    test_kinematics_gyroscopic_moment_and_antiskid()
    test_preload_classes_um_uh()
    test_zero_speed_division_by_zero_guard()
    test_all_8_load_cases_evaluation()
    test_ui_brevity_and_cleanliness()
    print("ALL BEARING SIZER UNIT TESTS PASSED SUCCESSFULLY!")
