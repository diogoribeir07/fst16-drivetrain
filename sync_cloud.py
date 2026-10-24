"""
sync_cloud.py
Bidirectional synchronization between local Excel workbooks (FST16_*.xlsx)
and static JSON files (data/*.json) for GitHub Pages 24/7 web hosting.
"""

import os
import sys
import json
import io
from datetime import datetime
import pandas as pd

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")

EXCEL_PARTS = os.path.join(BASE_DIR, "FST16_DT_Parts_Export.xlsx")
EXCEL_ASSEMBLIES = os.path.join(BASE_DIR, "FST16_Assemblies_Export_DL.xlsx")
EXCEL_REQUIREMENTS = os.path.join(BASE_DIR, "FST16_Requirements.xlsx")
EXCEL_TRASH = os.path.join(BASE_DIR, "FST16_Trash_Internal.xlsx")

JSON_PARTS = os.path.join(DATA_DIR, "parts.json")
JSON_ASSEMBLIES = os.path.join(DATA_DIR, "assemblies.json")
JSON_REQUIREMENTS = os.path.join(DATA_DIR, "requirements.json")
JSON_TRASH = os.path.join(DATA_DIR, "trash.json")
JSON_META = os.path.join(DATA_DIR, "meta.json")

def save_formatted_excel(df, filename, table_name="DataTable"):
    """Write dataframe to Excel with openpyxl styling without importing Streamlit."""
    try:
        if "Subassembly" in df.columns and "Part Name" in df.columns:
            df = df.sort_values(by=["Subassembly", "Part Name"]).reset_index(drop=True)

        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            sheet_name = 'Data'
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]

            for idx, col in enumerate(df.columns, 1):
                col_letter = get_column_letter(idx)

                if col in ["Description", "Specifications", "Constraints", "Part Configuration", "Image Path", 
                            "Stiffness Data", "LC Images", "Tolerancing Data", "Justification",
                            "Requirement", "Explanation", "Validation", "Related Parts", "Verification"]:
                    worksheet.column_dimensions[col_letter].width = 45
                    for row in range(2, len(df) + 2):
                        worksheet.cell(row=row, column=idx).alignment = Alignment(wrap_text=True, vertical='top')
                else:
                    col_lens = [len(str(x)) for x in df[col].dropna()]
                    max_len = max(max(col_lens) if col_lens else 0, len(str(col)))
                    worksheet.column_dimensions[col_letter].width = max(min(max_len + 5, 25), 14)
                    for row in range(2, len(df) + 2):
                        worksheet.cell(row=row, column=idx).alignment = Alignment(horizontal="center", vertical='top')

            if not df.empty:
                max_row = len(df) + 1
                max_col = len(df.columns)
                clean_table_name = table_name.replace(" ", "_").replace("-", "_")
                tab = Table(displayName=clean_table_name, ref=f"A1:{get_column_letter(max_col)}{max_row}")
                style = TableStyleInfo(
                    name="TableStyleMedium2", showFirstColumn=False,
                    showLastColumn=False, showRowStripes=True, showColumnStripes=False
                )
                tab.tableStyleInfo = style
                worksheet.add_table(tab)
        return True
    except PermissionError:
        print(f"  [ERROR] Close '{os.path.basename(filename)}' in Microsoft Excel before saving.")
        return False

def ensure_data_dir():
    if not os.path.exists(DATA_DIR):
        os.makedirs(DATA_DIR, exist_ok=True)

def read_excel_safely(filepath):
    """Read Excel workbook safely without locking files open in Microsoft Excel."""
    try:
        with open(filepath, "rb") as f:
            buf = io.BytesIO(f.read())
        return pd.read_excel(buf)
    except PermissionError:
        if os.name == "nt":
            import ctypes
            import msvcrt
            GENERIC_READ = 0x80000000
            FILE_SHARE_ALL = 0x00000001 | 0x00000002 | 0x00000004
            OPEN_EXISTING = 3
            FILE_ATTRIBUTE_NORMAL = 0x80
            handle = ctypes.windll.kernel32.CreateFileW(
                str(filepath), GENERIC_READ, FILE_SHARE_ALL, None, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL, None
            )
            if handle not in (-1, 0xFFFFFFFF, 0xFFFFFFFFFFFFFFFF):
                fd = msvcrt.open_osfhandle(handle, os.O_RDONLY | os.O_BINARY)
                with os.fdopen(fd, "rb") as f:
                    buf = io.BytesIO(f.read())
                return pd.read_excel(buf)
        raise

def sanitize_records_for_json(df):
    """Convert dataframe records into JSON-serializable list of dictionaries."""
    records = df.to_dict(orient="records")
    clean = []
    for row in records:
        clean_row = {}
        for k, v in row.items():
            if pd.isna(v) or v is None:
                clean_row[k] = None
            elif isinstance(v, (pd.Timestamp, datetime)):
                clean_row[k] = v.strftime("%Y-%m-%d %H:%M:%S")
            elif isinstance(v, float) and (v != v):  # NaN check
                clean_row[k] = None
            elif k == "Image Path" and isinstance(v, str):
                s_v = v.strip()
                if s_v.startswith("[") and s_v.endswith("]"):
                    try:
                        clean_row[k] = json.loads(s_v)
                    except Exception:
                        clean_row[k] = s_v
                else:
                    clean_row[k] = s_v
            else:
                clean_row[k] = v
        clean.append(clean_row)
    return clean

def prepare_df_for_excel(data):
    """Prepare records dictionary list for Excel dataframe writing, encoding lists/dicts cleanly."""
    clean_data = []
    for row in data:
        r = dict(row)
        for col in ["Image Path", "Stiffness Data", "LC Images", "Tolerancing Data"]:
            if col in r and isinstance(r[col], (list, dict)):
                r[col] = json.dumps(r[col], ensure_ascii=False)
        clean_data.append(r)
    return pd.DataFrame(clean_data)

def export_excel_to_json():
    """Export local Excel sheets to data/*.json for web consumption."""
    ensure_data_dir()
    print("=" * 60)
    print("Exporting Excel Workbooks -> data/*.json")
    print("=" * 60)

    # 1. Parts
    if os.path.exists(EXCEL_PARTS):
        df_parts = read_excel_safely(EXCEL_PARTS)
        records_parts = sanitize_records_for_json(df_parts)
        with open(JSON_PARTS, "w", encoding="utf-8") as f:
            json.dump(records_parts, f, indent=2, ensure_ascii=False)
        print(f"  [OK] Exported {len(records_parts)} part iterations -> data/parts.json")
    else:
        print(f"  [WARN] File not found: {EXCEL_PARTS}")

    # 2. Assemblies
    if os.path.exists(EXCEL_ASSEMBLIES):
        df_ass = read_excel_safely(EXCEL_ASSEMBLIES)
        records_ass = sanitize_records_for_json(df_ass)
        with open(JSON_ASSEMBLIES, "w", encoding="utf-8") as f:
            json.dump(records_ass, f, indent=2, ensure_ascii=False)
        print(f"  [OK] Exported {len(records_ass)} assemblies -> data/assemblies.json")
    else:
        print(f"  [WARN] File not found: {EXCEL_ASSEMBLIES}")

    # 3. Requirements
    if os.path.exists(EXCEL_REQUIREMENTS):
        df_req = read_excel_safely(EXCEL_REQUIREMENTS)
        records_req = sanitize_records_for_json(df_req)
        with open(JSON_REQUIREMENTS, "w", encoding="utf-8") as f:
            json.dump(records_req, f, indent=2, ensure_ascii=False)
        print(f"  [OK] Exported {len(records_req)} requirements -> data/requirements.json")
    else:
        print(f"  [WARN] File not found: {EXCEL_REQUIREMENTS}")

    # 4. Trash
    if os.path.exists(EXCEL_TRASH):
        df_trash = read_excel_safely(EXCEL_TRASH)
        records_trash = sanitize_records_for_json(df_trash)
        with open(JSON_TRASH, "w", encoding="utf-8") as f:
            json.dump(records_trash, f, indent=2, ensure_ascii=False)
        print(f"  [OK] Exported {len(records_trash)} trash items -> data/trash.json")
    else:
        print(f"  [WARN] File not found: {EXCEL_TRASH}")

    # Meta
    meta = {
        "last_sync": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source": "sync_cloud.py export",
        "version": "1.0.0"
    }
    with open(JSON_META, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    print("-" * 60)
    print("Export completed successfully.")
    print("=" * 60)

def import_json_to_excel():
    """Import cloud JSON updates into local Excel files."""
    ensure_data_dir()
    print("=" * 60)
    print("Importing data/*.json -> Local Excel Workbooks")
    print("=" * 60)

    # 1. Parts
    if os.path.exists(JSON_PARTS):
        with open(JSON_PARTS, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list) and len(data) > 0:
            df = prepare_df_for_excel(data)
            ok = save_formatted_excel(df, EXCEL_PARTS, table_name="PartsTable")
            if ok:
                print(f"  [OK] Saved {len(df)} parts into {os.path.basename(EXCEL_PARTS)}")
            else:
                print(f"  [FAIL] Could not save {os.path.basename(EXCEL_PARTS)}")

    # 2. Assemblies
    if os.path.exists(JSON_ASSEMBLIES):
        with open(JSON_ASSEMBLIES, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list) and len(data) > 0:
            df = prepare_df_for_excel(data)
            ok = save_formatted_excel(df, EXCEL_ASSEMBLIES, table_name="AssembliesTable")
            if ok:
                print(f"  [OK] Saved {len(df)} assemblies into {os.path.basename(EXCEL_ASSEMBLIES)}")
            else:
                print(f"  [FAIL] Could not save {os.path.basename(EXCEL_ASSEMBLIES)}")

    # 3. Requirements
    if os.path.exists(JSON_REQUIREMENTS):
        with open(JSON_REQUIREMENTS, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list) and len(data) > 0:
            df = prepare_df_for_excel(data)
            ok = save_formatted_excel(df, EXCEL_REQUIREMENTS, table_name="RequirementsTable")
            if ok:
                print(f"  [OK] Saved {len(df)} requirements into {os.path.basename(EXCEL_REQUIREMENTS)}")
            else:
                print(f"  [FAIL] Could not save {os.path.basename(EXCEL_REQUIREMENTS)}")

    # 4. Trash
    if os.path.exists(JSON_TRASH):
        with open(JSON_TRASH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            df = prepare_df_for_excel(data)
            ok = save_formatted_excel(df, EXCEL_TRASH, table_name="TrashTable")
            if ok:
                print(f"  [OK] Saved {len(df)} trash items into {os.path.basename(EXCEL_TRASH)}")
            else:
                print(f"  [FAIL] Could not save {os.path.basename(EXCEL_TRASH)}")

    print("-" * 60)
    print("Import completed successfully.")
    print("=" * 60)

if __name__ == "__main__":
    action = sys.argv[1].lower() if len(sys.argv) > 1 else "export"
    if action == "import":
        import_json_to_excel()
    else:
        export_excel_to_json()
